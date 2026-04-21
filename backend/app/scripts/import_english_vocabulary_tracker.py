from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.shared.db import SessionLocal
from app.features.topics.model import Topic
from app.features.words.model import Word
from app.features.stats.model import WordProgressEvent  # noqa: F401 — registers the ORM class so Word.progress_events resolves
from app.features.words.model import word_topics
from app.scripts.xlsx_mapping import (
    IGNORED_SHEETS,
    SheetConfig,
    is_empty_row,
    is_repeated_header_row,
    normalize_headers,
    normalize_text,
    parse_knowledge_level,
    resolve_sheet_config,
    row_to_dict,
    slugify,
)


def get_or_create_topic(db: Session, sheet_name: str) -> tuple[Topic, bool]:
    slug = slugify(sheet_name)
    topic: Topic | None = db.scalar(select(Topic).where(Topic.slug == slug))
    if topic:
        assert topic is not None
        return topic, False
    new_topic = Topic(name=sheet_name, slug=slug, description=None, is_active=True)
    db.add(new_topic)
    db.flush()
    return new_topic, True


def existing_terms(db: Session, topic: Topic) -> set[str]:
    rows = db.scalars(
        select(Word.term).join(word_topics, Word.id == word_topics.c.word_id).where(word_topics.c.topic_id == topic.id)
    ).all()
    return {t.lower() for t in rows}


def build_word(topic: Topic, term: str, translations: str, config: SheetConfig, row_data: dict, sheet_name: str, row_number: int) -> Word:
    return Word(
        topics=[topic],
        term=term,
        translations=translations,
        part_of_speech=config.part_of_speech,
        knowledge_level=parse_knowledge_level(row_data.get("Knowledge"), sheet_name=sheet_name, row_number=row_number),
        past_simple=normalize_text(row_data.get(config.past_simple_header) if config.past_simple_header else None),
        past_participle=normalize_text(row_data.get(config.past_participle_header) if config.past_participle_header else None),
        countability=normalize_text(row_data.get(config.countability_header) if config.countability_header else None),
        pattern=normalize_text(row_data.get(config.pattern_header) if config.pattern_header else None),
        example=normalize_text(row_data.get(config.example_header) if config.example_header else None),
        notes=normalize_text(row_data.get(config.notes_header) if config.notes_header else None),
        is_active=True,
    )


def import_workbook(file_path: Path) -> None:
    if not file_path.exists():
        raise FileNotFoundError(f"Workbook not found: {file_path}")

    workbook = load_workbook(filename=file_path, data_only=True)
    topics_created = topics_existing = words_added = words_skipped = 0

    with SessionLocal() as db:
        try:
            for worksheet in workbook.worksheets:
                if worksheet.title in IGNORED_SHEETS:
                    continue

                header_row = tuple(cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1)))
                headers = normalize_headers(header_row)
                config = resolve_sheet_config(worksheet.title, headers)

                topic, created = get_or_create_topic(db, worksheet.title)
                if created:
                    topics_created += 1
                else:
                    topics_existing += 1

                seen = existing_terms(db, topic)

                for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
                    row_data = row_to_dict(headers, row)

                    if is_empty_row(row_data, config) or is_repeated_header_row(row_data):
                        continue

                    term = normalize_text(row_data.get(config.term_header))
                    translations = normalize_text(row_data.get(config.translations_header))
                    if term is None or translations is None:
                        continue

                    if term.lower() in seen:
                        words_skipped += 1
                        continue

                    db.add(build_word(topic, term, translations, config, row_data, worksheet.title, row_number))
                    seen.add(term.lower())
                    words_added += 1

            db.commit()
        except Exception:
            db.rollback()
            raise

    print(f"\nImport complete: {file_path}")
    print(f"  Topics  — created: {topics_created}, already existed: {topics_existing}")
    print(f"  Words   — added:   {words_added}, skipped (duplicates): {words_skipped}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Sync the English Vocabulary Tracker workbook into the Lexora database. "
                    "Safe to run multiple times — only new words are added, existing ones are never touched."
    )
    parser.add_argument("file", type=Path, help="Path to the .xlsx workbook")
    args = parser.parse_args()
    import_workbook(file_path=args.file)


if __name__ == "__main__":
    main()
