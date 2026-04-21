from __future__ import annotations

from typing import NamedTuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.features.topics.model import Topic
from app.features.words.model import Word

HEADERS: dict[str, tuple[str, ...]] = {
    "verb":         ("Knowledge", "Word", "Russian translations", "Typical prepositions / patterns", "Examples (EN + RU)"),
    "irregular_verb": ("Knowledge", "Base form", "Past Simple", "Past Participle", "Russian translations", "Typical prepositions / patterns", "Examples (EN + RU)"),
    "noun":         ("Knowledge", "Word", "Russian translations", "Countability"),
    "adjective":    ("Knowledge", "Word", "Russian translations"),
    "preposition":  ("Knowledge", "Word", "Russian translations"),
    "phrase":       ("Knowledge", "Phrase", "Russian translations", "Meaning / usage note"),
    "adverb":       ("Knowledge", "Word", "Russian translations", "Typical position / usage"),
}

_HEADER_FILL  = PatternFill("solid", fgColor="1E293B")
_HEADER_FONT  = Font(bold=True, color="FFFFFF", size=10)
_HEADER_ALIGN = Alignment(horizontal="left", vertical="center")
_COL_WIDTHS   = [12, 28, 40, 30, 60]


def detect_sheet_type(words: list[Word]) -> str:
    pos_set = {w.part_of_speech for w in words if w.part_of_speech}
    if "verb" in pos_set:
        return "irregular_verb" if any(w.past_simple for w in words) else "verb"
    for pos in ("noun", "adjective", "preposition", "phrase", "adverb"):
        if pos in pos_set:
            return pos
    return "noun"


class WordRow(NamedTuple):
    knowledge: int
    term: str
    translations: str | None
    extra_1: str | None = None
    extra_2: str | None = None
    extra_3: str | None = None
    extra_4: str | None = None


def word_to_row(word: Word, sheet_type: str) -> WordRow:
    level = word.knowledge_level or 1
    if sheet_type == "verb":
        return WordRow(level, word.term, word.translations, word.pattern, word.example)
    if sheet_type == "irregular_verb":
        return WordRow(level, word.term, word.past_simple, word.past_participle, word.translations, word.pattern, word.example)
    if sheet_type == "noun":
        return WordRow(level, word.term, word.translations, word.countability)
    if sheet_type in ("adjective", "preposition"):
        return WordRow(level, word.term, word.translations)
    if sheet_type == "phrase":
        return WordRow(level, word.term, word.translations, word.notes)
    if sheet_type == "adverb":
        return WordRow(level, word.term, word.translations, word.pattern)
    return WordRow(level, word.term, word.translations)


def write_sheet(wb: Workbook, topic: Topic, words: list[Word]) -> None:
    sheet_type = detect_sheet_type(words)
    headers = HEADERS[sheet_type]
    ws = wb.create_sheet(title=topic.name[:31])

    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN

    for word in sorted(words, key=lambda w: (w.knowledge_level or 99, w.term.lower())):
        ws.append(word_to_row(word, sheet_type))

    for i, width in enumerate(_COL_WIDTHS[:len(headers)], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
