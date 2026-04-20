from __future__ import annotations

import re
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.features.stats.service import record_level_change
from app.features.topics.model import Topic
from app.features.topics.service import TopicSlugConflictError, assert_slug_available
from app.features.words.domain import assert_no_duplicate_word, existing_normalized_terms
from app.features.words.model import Word
from app.features.words.repository import get_word_by_id_including_deleted, sync_word_multivalue_fields
from app.features.words.schemas import WorkbookImportResponse, WorkbookImportSheetSummary
from app.shared.constraints import TOPIC_SLUG_MAX_LEN
from app.shared.text import normalize_term, slugify

LISTS_SHEET_NAME = "Lists"
META_SHEET_NAME = "__lexora_meta"

EXPORT_COLUMNS = [
    ("knowledge_level", "Knowledge"),
    ("term", "Word"),
    ("translations", "Russian translations"),
    ("pattern", "Typical prepositions / patterns"),
    ("examples", "Examples (EN + RU)"),
    ("countability", "Countability"),
    ("part_of_speech", "Part of speech"),
    ("past_simple", "Past simple"),
    ("past_participle", "Past participle"),
    ("notes", "Notes"),
    ("word_id", "Word ID"),
]

HEADER_ALIASES: dict[str, set[str]] = {
    "knowledge_level": {"knowledge", "level"},
    "term": {"word", "term", "base form", "phrase"},
    "translations": {
        "russian translations",
        "translations",
        "translation",
        "russian translation",
    },
    "pattern": {
        "typical prepositions / patterns",
        "pattern",
        "patterns",
        "prepositions / patterns",
        "typical position / usage",
    },
    "examples": {"examples (en + ru)", "examples", "example", "example (en+ru)"},
    "countability": {"countability"},
    "part_of_speech": {"part of speech", "pos"},
    "past_simple": {"past simple"},
    "past_participle": {"past participle"},
    "notes": {"notes", "note", "meaning / usage note"},
    "word_id": {"word id", "id"},
}

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)

LEVEL_FILLS = {
    1: PatternFill(fill_type="solid", fgColor="FEE2E2"),
    2: PatternFill(fill_type="solid", fgColor="FEF3C7"),
    3: PatternFill(fill_type="solid", fgColor="DBEAFE"),
    4: PatternFill(fill_type="solid", fgColor="DCFCE7"),
    5: PatternFill(fill_type="solid", fgColor="E5E7EB"),
}

PART_OF_SPEECH_VALUES = [
    "noun",
    "verb",
    "adjective",
    "adverb",
    "phrase",
    "preposition",
    "other",
]
COUNTABILITY_VALUES = ["Countable", "Uncountable", "Both"]
EXTRA_EMPTY_ROWS = 30


class InvalidWorkbookError(Exception):
    pass


def _safe_sheet_title(raw_title: str, used_titles: set[str]) -> str:
    cleaned = re.sub(r"[:\\/?*\[\]]", " ", raw_title).strip()
    if not cleaned:
        cleaned = "Topic"
    cleaned = re.sub(r"\s+", " ", cleaned)

    base = cleaned[:31]
    candidate = base
    suffix = 2
    while candidate in used_titles:
        suffix_text = f" ({suffix})"
        candidate = f"{base[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1

    used_titles.add(candidate)
    return candidate


def _set_column_widths(ws) -> None:
    widths = {
        "A": 12,
        "B": 28,
        "C": 36,
        "D": 28,
        "E": 42,
        "F": 16,
        "G": 18,
        "H": 18,
        "I": 20,
        "J": 28,
        "K": 12,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.column_dimensions["K"].hidden = True


def _apply_base_styling(ws, last_row: int) -> None:
    for row in ws.iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=len(EXPORT_COLUMNS)):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{last_row}"


def _add_dynamic_row_colors(ws, last_row: int) -> None:
    if last_row < 2:
        return

    for level, fill in LEVEL_FILLS.items():
        ws.conditional_formatting.add(
            f"A2:J{last_row}",
            FormulaRule(formula=[f"$A2={level}"], stopIfTrue=False, fill=fill),
        )


def _add_validations(workbook: Workbook, ws, last_row: int) -> None:
    knowledge_dv = DataValidation(
        type="list", formula1=f"={LISTS_SHEET_NAME}!$A$1:$A$5", allow_blank=True
    )
    countability_dv = DataValidation(
        type="list", formula1=f"={LISTS_SHEET_NAME}!$B$1:$B$3", allow_blank=True
    )
    pos_dv = DataValidation(
        type="list", formula1=f"={LISTS_SHEET_NAME}!$C$1:$C$7", allow_blank=True
    )

    ws.add_data_validation(knowledge_dv)
    ws.add_data_validation(countability_dv)
    ws.add_data_validation(pos_dv)

    knowledge_dv.add(f"A2:A{last_row}")
    countability_dv.add(f"F2:F{last_row}")
    pos_dv.add(f"G2:G{last_row}")


def _create_lists_sheet(workbook: Workbook) -> None:
    ws = workbook.create_sheet(LISTS_SHEET_NAME)
    for idx, value in enumerate(["1", "2", "3", "4", "5"], start=1):
        ws.cell(row=idx, column=1, value=value)
    for idx, value in enumerate(COUNTABILITY_VALUES, start=1):
        ws.cell(row=idx, column=2, value=value)
    for idx, value in enumerate(PART_OF_SPEECH_VALUES, start=1):
        ws.cell(row=idx, column=3, value=value)
    ws.sheet_state = "hidden"


def _create_meta_sheet(workbook: Workbook, mappings: list[tuple[str, int, str]]) -> None:
    ws = workbook.create_sheet(META_SHEET_NAME)
    ws.append(["sheet_name", "topic_id", "topic_name", "exported_at"])
    exported_at = datetime.now(UTC).isoformat()
    for sheet_name, topic_id, topic_name in mappings:
        ws.append([sheet_name, topic_id, topic_name, exported_at])
    ws.sheet_state = "hidden"


def _translation_entries_for_export(word: Word) -> str:
    if word.translation_items:
        return "\n".join(item.value for item in word.translation_items)
    return word.translations


def _example_entries_for_export(word: Word) -> str | None:
    if word.example_items:
        return "\n".join(item.value for item in word.example_items)
    return word.example


def build_words_workbook(db: Session) -> tuple[str, bytes]:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    _create_lists_sheet(workbook)

    topics = list(
        db.scalars(select(Topic).where(Topic.deleted_at.is_(None)).order_by(Topic.name.asc())).all()
    )

    used_titles: set[str] = set()
    meta_mappings: list[tuple[str, int, str]] = []

    if not topics:
        ws = workbook.create_sheet("Vocabulary")
        for idx, (_, header) in enumerate(EXPORT_COLUMNS, start=1):
            ws.cell(row=1, column=idx, value=header)
        _set_column_widths(ws)
        _apply_base_styling(ws, 2)
        _add_dynamic_row_colors(ws, 2)
        _add_validations(workbook, ws, 2)
    else:
        for topic in topics:
            words = list(
                db.scalars(
                    select(Word)
                    .options(
                        selectinload(Word.translation_items),
                        selectinload(Word.example_items),
                        selectinload(Word.topics),
                    )
                    .where(Word.deleted_at.is_(None))
                    .where(Word.topics.any((Topic.id == topic.id) & Topic.deleted_at.is_(None)))
                    .order_by(Word.term.asc())
                ).all()
            )

            sheet_title = _safe_sheet_title(topic.name, used_titles)
            meta_mappings.append((sheet_title, topic.id, topic.name))
            ws = workbook.create_sheet(sheet_title)

            for idx, (_, header) in enumerate(EXPORT_COLUMNS, start=1):
                ws.cell(row=1, column=idx, value=header)

            row_idx = 2
            for word in words:
                ws.cell(row=row_idx, column=1, value=word.knowledge_level)
                ws.cell(row=row_idx, column=2, value=word.term)
                ws.cell(row=row_idx, column=3, value=_translation_entries_for_export(word))
                ws.cell(row=row_idx, column=4, value=word.pattern)
                ws.cell(row=row_idx, column=5, value=_example_entries_for_export(word))
                ws.cell(row=row_idx, column=6, value=word.countability)
                ws.cell(row=row_idx, column=7, value=word.part_of_speech)
                ws.cell(row=row_idx, column=8, value=word.past_simple)
                ws.cell(row=row_idx, column=9, value=word.past_participle)
                ws.cell(row=row_idx, column=10, value=word.notes)
                ws.cell(row=row_idx, column=11, value=word.id)
                row_idx += 1

            last_row = max(2, row_idx + EXTRA_EMPTY_ROWS - 1)
            _set_column_widths(ws)
            _apply_base_styling(ws, last_row)
            _add_dynamic_row_colors(ws, last_row)
            _add_validations(workbook, ws, last_row)

    _create_meta_sheet(workbook, meta_mappings)

    output = BytesIO()
    workbook.save(output)
    filename = f"lexora-vocabulary-{datetime.now(UTC).strftime('%Y%m%d-%H%M%S')}.xlsx"
    return filename, output.getvalue()


def _normalize_header(value: str | None) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", value.strip().lower())


def _build_header_map(ws) -> dict[str, int]:
    result: dict[str, int] = {}
    for column_idx in range(1, ws.max_column + 1):
        raw = ws.cell(row=1, column=column_idx).value
        header = _normalize_header(str(raw) if raw is not None else "")
        for canonical, aliases in HEADER_ALIASES.items():
            if header in aliases:
                result[canonical] = column_idx
                break
    return result


def _read_str(ws, row_idx: int, column_idx: int | None) -> str | None:
    if column_idx is None:
        return None
    value = ws.cell(row=row_idx, column=column_idx).value
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _read_optional_int(
    ws, row_idx: int, column_idx: int | None, label: str, sheet_name: str
) -> int | None:
    if column_idx is None:
        return None
    raw = ws.cell(row=row_idx, column=column_idx).value
    if raw is None or str(raw).strip() == "":
        return None
    try:
        value = int(raw)
    except (TypeError, ValueError) as exc:
        raise InvalidWorkbookError(
            f"{sheet_name}, row {row_idx}: invalid {label} value '{raw}'"
        ) from exc
    return value


def _split_translation_cell(value: str | None) -> list[str]:
    if value is None:
        return []
    text = value.strip()
    if not text:
        return []
    if "\n" in text or ";" in text:
        parts = re.split(r"(?:\r?\n|;)+", text)
        return [part.strip() for part in parts if part and part.strip()]
    return [text]


def _split_examples_cell(value: str | None) -> list[str]:
    if value is None:
        return []
    text = value.strip()
    if not text:
        return []
    parts = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    return [part.strip() for part in parts if part and part.strip()]


def _normalize_countability(value: str | None) -> str | None:
    if value is None:
        return None
    normalized = value.strip()
    return normalized or None


def _normalize_part_of_speech(
    raw_value: str | None,
    countability: str | None,
    past_simple: str | None,
    past_participle: str | None,
    existing_value: str | None = None,
) -> str | None:
    if raw_value and raw_value.strip():
        return raw_value.strip().lower()
    if existing_value:
        return existing_value
    if past_simple or past_participle:
        return "verb"
    if countability:
        return "noun"
    return None


def _get_or_create_topic(db: Session, topic_name: str) -> Topic:
    topic = db.scalar(select(Topic).where(Topic.name == topic_name).where(Topic.deleted_at.is_(None)))
    if topic is not None:
        return topic

    deleted = db.scalar(select(Topic).where(Topic.name == topic_name).where(Topic.deleted_at.isnot(None)))
    if deleted is not None:
        raise InvalidWorkbookError(
            f"Workbook references topic '{topic_name}', but that topic is currently in Trash."
        )

    slug = slugify(topic_name, max_len=TOPIC_SLUG_MAX_LEN)
    if not slug:
        raise InvalidWorkbookError(f"Cannot generate a valid slug for topic '{topic_name}'")
    try:
        assert_slug_available(db, slug)
    except TopicSlugConflictError as exc:
        raise InvalidWorkbookError(str(exc)) from exc

    topic = Topic(name=topic_name, slug=slug, description=None, is_active=True)
    db.add(topic)
    db.flush()
    return topic


def _find_existing_word(db: Session, topic_id: int, word_id: int | None, term: str) -> Word | None:
    if word_id is not None:
        word = get_word_by_id_including_deleted(db, word_id)
        if word is None:
            raise InvalidWorkbookError(f"Workbook references unknown word id {word_id}")
        if word.deleted_at is not None:
            raise InvalidWorkbookError(f"Workbook references deleted word id {word_id}")
        return word

    candidates = list(
        db.scalars(
            select(Word)
            .options(
                selectinload(Word.topics),
                selectinload(Word.translation_items),
                selectinload(Word.example_items),
            )
            .where(Word.deleted_at.is_(None))
            .where(Word.topics.any(Topic.id == topic_id))
        ).all()
    )

    normalized = normalize_term(term)
    for candidate in candidates:
        if normalize_term(candidate.term) == normalized:
            return candidate
    return None


def _read_meta_topic_names(workbook) -> dict[str, str]:
    if META_SHEET_NAME not in workbook.sheetnames:
        return {}

    ws = workbook[META_SHEET_NAME]
    result: dict[str, str] = {}
    for row_idx in range(2, ws.max_row + 1):
        sheet_name = _read_str(ws, row_idx, 1)
        topic_name = _read_str(ws, row_idx, 3)
        if sheet_name and topic_name:
            result[sheet_name] = topic_name
    return result


def _import_sheet(
    db: Session, ws, header_map: dict[str, int], topic: Topic
) -> WorkbookImportSheetSummary:
    created = 0
    updated = 0
    skipped = 0

    has_knowledge = "knowledge_level" in header_map
    has_pattern = "pattern" in header_map
    has_examples = "examples" in header_map
    has_countability = "countability" in header_map
    has_part_of_speech = "part_of_speech" in header_map
    has_past_simple = "past_simple" in header_map
    has_past_participle = "past_participle" in header_map
    has_notes = "notes" in header_map

    for row_idx in range(2, ws.max_row + 1):
        term = _read_str(ws, row_idx, header_map.get("term"))
        if not term:
            continue

        translations_text = _read_str(ws, row_idx, header_map.get("translations"))
        if not translations_text:
            raise InvalidWorkbookError(
                f"{ws.title}, row {row_idx}: translations are required for '{term}'"
            )

        word_id = _read_optional_int(ws, row_idx, header_map.get("word_id"), "word id", ws.title)
        knowledge_value = (
            _read_optional_int(ws, row_idx, header_map.get("knowledge_level"), "knowledge", ws.title)
            if has_knowledge
            else None
        )
        if knowledge_value is not None and knowledge_value not in {1, 2, 3, 4, 5}:
            raise InvalidWorkbookError(f"{ws.title}, row {row_idx}: knowledge must be between 1 and 5")

        pattern = _read_str(ws, row_idx, header_map.get("pattern")) if has_pattern else None
        examples_text = _read_str(ws, row_idx, header_map.get("examples")) if has_examples else None
        countability = (
            _normalize_countability(_read_str(ws, row_idx, header_map.get("countability")))
            if has_countability
            else None
        )
        past_simple = _read_str(ws, row_idx, header_map.get("past_simple")) if has_past_simple else None
        past_participle = (
            _read_str(ws, row_idx, header_map.get("past_participle"))
            if has_past_participle
            else None
        )
        notes = _read_str(ws, row_idx, header_map.get("notes")) if has_notes else None

        existing = _find_existing_word(db, topic.id, word_id, term)

        if existing is not None:
            if topic.id not in {t.id for t in existing.topics}:
                existing.topics.append(topic)

            assert_no_duplicate_word(
                term,
                existing_normalized_terms(db, [topic.id], exclude_word_id=existing.id),
            )

            old_level = existing.knowledge_level

            existing.term = term
            existing.translations = translations_text
            existing.knowledge_level = knowledge_value if has_knowledge else existing.knowledge_level
            existing.countability = countability if has_countability else existing.countability
            existing.pattern = pattern if has_pattern else existing.pattern
            existing.notes = notes if has_notes else existing.notes
            existing.past_simple = past_simple if has_past_simple else existing.past_simple
            existing.past_participle = (
                past_participle if has_past_participle else existing.past_participle
            )
            existing.part_of_speech = _normalize_part_of_speech(
                _read_str(ws, row_idx, header_map.get("part_of_speech"))
                if has_part_of_speech
                else None,
                existing.countability,
                existing.past_simple,
                existing.past_participle,
                existing.part_of_speech,
            )

            sync_word_multivalue_fields(
                existing,
                translations_text,
                _split_translation_cell(translations_text),
                examples_text if has_examples else existing.example,
                _split_examples_cell(examples_text)
                if has_examples
                else [item.value for item in existing.example_items],
            )

            if existing.knowledge_level != old_level and existing.knowledge_level is not None:
                record_level_change(db, existing.id, old_level, existing.knowledge_level, "xlsx_import")

            db.add(existing)
            updated += 1
            continue

        assert_no_duplicate_word(term, existing_normalized_terms(db, [topic.id]))

        knowledge_for_create = knowledge_value if knowledge_value is not None else 1
        part_of_speech = _normalize_part_of_speech(
            _read_str(ws, row_idx, header_map.get("part_of_speech"))
            if has_part_of_speech
            else None,
            countability,
            past_simple,
            past_participle,
            None,
        )

        word = Word(
            term=term,
            past_simple=past_simple,
            past_participle=past_participle,
            translations=translations_text,
            part_of_speech=part_of_speech,
            knowledge_level=knowledge_for_create,
            countability=countability,
            pattern=pattern,
            example=examples_text,
            notes=notes,
            is_active=True,
            topics=[topic],
        )
        sync_word_multivalue_fields(
            word,
            translations_text,
            _split_translation_cell(translations_text),
            examples_text,
            _split_examples_cell(examples_text),
        )

        db.add(word)
        created += 1

    return WorkbookImportSheetSummary(
        sheet_name=ws.title,
        topic_name=topic.name,
        created=created,
        updated=updated,
        skipped=skipped,
    )


def import_words_workbook(db: Session, content: bytes) -> WorkbookImportResponse:
    try:
        workbook = load_workbook(filename=BytesIO(content))
    except Exception as exc:
        raise InvalidWorkbookError("Failed to read XLSX workbook") from exc

    meta_topic_names = _read_meta_topic_names(workbook)

    summaries: list[WorkbookImportSheetSummary] = []
    total_created = 0
    total_updated = 0
    total_skipped = 0

    for ws in workbook.worksheets:
        if ws.sheet_state != "visible":
            continue
        if ws.title in {LISTS_SHEET_NAME, META_SHEET_NAME}:
            continue

        header_map = _build_header_map(ws)
        if "term" not in header_map or "translations" not in header_map:
            continue

        topic_name = meta_topic_names.get(ws.title, ws.title)
        topic = _get_or_create_topic(db, topic_name)

        summary = _import_sheet(db, ws, header_map, topic)
        summaries.append(summary)
        total_created += summary.created
        total_updated += summary.updated
        total_skipped += summary.skipped

    if not summaries:
        raise InvalidWorkbookError("No importable sheets found in workbook")

    db.commit()
    return WorkbookImportResponse(
        created=total_created,
        updated=total_updated,
        skipped=total_skipped,
        sheets=summaries,
    )
