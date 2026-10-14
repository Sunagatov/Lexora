"""Read content updates by Word ID without interpreting learning progress."""

from __future__ import annotations

from dataclasses import dataclass, field
from hashlib import sha256
from io import BytesIO

from openpyxl import load_workbook

from app.features.words.constants import CEFR_LEVELS, PART_OF_SPEECH_VALUES, REGISTER_VALUES
from app.features.words.workbook.cells import (
    _build_header_map,
    _split_translation_cell,
)
from app.features.words.workbook.cells import (
    _ordered_unique as ordered_unique,
)
from app.features.words.workbook.cells import (
    _split_examples_cell as parse_examples,
)
from app.features.words.workbook.format import (
    COUNTABILITY_VALUES,
    LISTS_SHEET_NAME,
    MAX_WORKBOOK_BYTES,
    META_SHEET_NAME,
    InvalidWorkbookError,
)

# The app has a deliberately small POS taxonomy. Keep the original finer label in Notes.
POS_ALIASES = {
    "proper noun": "noun",
    "prepositional verb": "verb",
    "abbreviation": "other",
    "fixed expression": "phrase",
    "verb phrase": "phrase",
    "noun phrase / fixed expression": "phrase",
    "prepositional / adverbial phrase": "phrase",
    "question / fixed expression": "phrase",
    "adjectival / adverbial phrase": "phrase",
    "date format": "phrase",
    "time expression": "phrase",
    "time / adverbial expression": "phrase",
}
SCALAR_FIELDS = (
    "term",
    "definition",
    "pattern",
    "countability",
    "part_of_speech",
    "cefr_level",
    "register",
)


def positive_id(value: object) -> int:
    if isinstance(value, bool) or value is None:
        raise ValueError("Expected a positive integer ID")
    try:
        number = int(str(value))
    except (ValueError, TypeError):
        if not isinstance(value, float) or not value.is_integer():
            raise ValueError("Expected a positive integer ID") from None
        number = int(value)
    if number <= 0:
        raise ValueError("Expected a positive integer ID")
    return number


@dataclass
class SyncSource:
    workbook_sha256: str
    rows: int = 0
    sheets: int = 0
    records: dict[int, dict] = field(default_factory=dict)
    conflicts: list[dict] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    normalizations: list[dict] = field(default_factory=list)
    topic_annotations: list[dict] = field(default_factory=list)


def read_sync_workbook(content: bytes, resolutions: dict | None = None) -> SyncSource:
    if len(content) > MAX_WORKBOOK_BYTES:
        raise InvalidWorkbookError("Workbook exceeds the 10 MB limit")
    try:
        workbook = load_workbook(BytesIO(content), keep_links=False)
    except Exception as exc:
        raise InvalidWorkbookError("Failed to read XLSX workbook") from exc
    source = SyncSource(sha256(content).hexdigest())
    resolutions = {} if resolutions is None else resolutions
    if not isinstance(resolutions, dict) or any(
        not isinstance(item, dict) for item in resolutions.values()
    ):
        raise InvalidWorkbookError("Resolutions must map Word IDs to field/value objects")
    variants: dict[int, dict[str, list[dict]]] = {}
    try:
        if META_SHEET_NAME not in workbook:
            raise InvalidWorkbookError(
                "Synchronization requires exported topic IDs in __lexora_meta"
            )
        topic_ids: dict[str, int] = {}
        for values in workbook[META_SHEET_NAME].iter_rows(min_row=2, values_only=True):
            if not values or not values[0]:
                continue
            name = str(values[0])
            if name in topic_ids:
                raise InvalidWorkbookError(f"Duplicate topic metadata for {name}")
            try:
                topic_ids[name] = positive_id(values[1])
            except ValueError as exc:
                raise InvalidWorkbookError(f"{name}: invalid topic ID") from exc
        for ws in workbook:
            if ws.title in {LISTS_SHEET_NAME, META_SHEET_NAME} or ws.sheet_state != "visible":
                continue
            header = _build_header_map(ws)
            if "term" not in header:
                continue
            source.sheets += 1
            if ws.title not in topic_ids or "word_id" not in header:
                source.errors.append(
                    {"sheet": ws.title, "error": "Missing exported topic ID or Word ID column"}
                )
                continue
            for row_number in range(2, ws.max_row + 1):
                values = {key: ws.cell(row_number, column).value for key, column in header.items()}
                if not any(value is not None and str(value).strip() for value in values.values()):
                    continue
                location = {"sheet": ws.title, "row": row_number}
                word_id = None
                try:
                    word_id = positive_id(values.get("word_id"))
                    if any(
                        ws.cell(row_number, column).data_type == "f" for column in header.values()
                    ):
                        raise ValueError("Formula cells are not accepted as content updates")
                    term = str(values.get("term") or "").strip()
                    if not term or len(term) > 255:
                        raise ValueError("Word must contain 1 to 255 characters")
                    record = source.records.setdefault(
                        word_id, {"topic_ids": set(), "locations": [], "values": {}}
                    )
                    record["topic_ids"].add(topic_ids[ws.title])
                    record["locations"].append(location)
                    source.rows += 1
                    parsed = {
                        key: str(values[key]).strip()
                        for key in SCALAR_FIELDS
                        if values.get(key) is not None and str(values[key]).strip()
                    }
                    raw_pos = parsed.get("part_of_speech", "").lower()
                    if raw_pos:
                        parsed["part_of_speech"] = POS_ALIASES.get(raw_pos, raw_pos)
                        if parsed["part_of_speech"] not in PART_OF_SPEECH_VALUES:
                            raise ValueError(f"Unsupported part of speech: {raw_pos}")
                        if raw_pos in POS_ALIASES:
                            source.normalizations.append(
                                {
                                    "word_id": word_id,
                                    **location,
                                    "from": raw_pos,
                                    "to": parsed["part_of_speech"],
                                }
                            )
                    for key, allowed in (
                        ("cefr_level", CEFR_LEVELS),
                        ("register", REGISTER_VALUES),
                        ("countability", COUNTABILITY_VALUES),
                    ):
                        if key in parsed and parsed[key] not in allowed:
                            raise ValueError(f"Invalid {key}: {parsed[key]}")
                    if parsed.get("countability") and parsed.get("part_of_speech") not in (
                        None,
                        "noun",
                    ):
                        raise ValueError("Countability is supported only for nouns")
                    if (
                        values.get("translations") is not None
                        and str(values["translations"]).strip()
                    ):
                        parsed["translations"] = ordered_unique(
                            _split_translation_cell(str(values["translations"]))
                        )
                    if values.get("examples") is not None and str(values["examples"]).strip():
                        parsed["examples"] = parse_examples(str(values["examples"]))
                    notes = str(values.get("notes") or "").strip()
                    if raw_pos in POS_ALIASES:
                        notes = "\n".join(
                            filter(None, [notes, f"Workbook part of speech: {raw_pos}"])
                        )
                    if notes:
                        record["values"]["notes"] = ordered_unique(
                            [*record["values"].get("notes", []), *notes.splitlines()]
                        )
                        source.topic_annotations.append(
                            {"word_id": word_id, **location, "notes": notes}
                        )
                    for key, value in parsed.items():
                        entries = variants.setdefault(word_id, {}).setdefault(key, [])
                        matching = next(
                            (entry for entry in entries if entry["value"] == value), None
                        )
                        if matching is None:
                            entries.append({"value": value, "locations": [location]})
                        else:
                            matching["locations"].append(location)
                except (ValueError, InvalidWorkbookError) as exc:
                    source.errors.append({"word_id": word_id, **location, "error": str(exc)})
        for word_id, fields in variants.items():
            overrides = resolutions.get(str(word_id), {})
            for name, entries in fields.items():
                if name in overrides:
                    # A resolution must explicitly choose an observed, already validated variant.
                    chosen = next(
                        (entry["value"] for entry in entries if entry["value"] == overrides[name]),
                        None,
                    )
                    if chosen is None:
                        source.errors.append(
                            {
                                "word_id": word_id,
                                "field": name,
                                "error": "Resolution is not an observed variant",
                            }
                        )
                    else:
                        source.records[word_id]["values"][name] = chosen
                elif len(entries) == 1:
                    source.records[word_id]["values"][name] = entries[0]["value"]
                else:
                    source.conflicts.append(
                        {"word_id": word_id, "field": name, "variants": entries}
                    )
        for key, overrides in resolutions.items():
            if (
                not str(key).isdigit()
                or int(key) not in variants
                or set(overrides) - set(variants[int(key)])
            ):
                source.errors.append(
                    {"word_id": key, "error": "Resolution references an unknown record or field"}
                )
        if not source.records and not source.errors:
            raise InvalidWorkbookError("No vocabulary records with Word IDs found")
        return source
    finally:
        workbook.close()
