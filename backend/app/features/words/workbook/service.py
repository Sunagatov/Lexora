from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, load_only, selectinload

from app.features.topics.model import Topic
from app.features.words.model import Word, word_topics
from app.features.words.workbook.cells import _normalize_countability
from app.features.words.workbook.format import (
    EXPORT_COLUMNS,
    EXTRA_EMPTY_ROWS,
    InvalidWorkbookError,
    _add_dynamic_row_colors,
    _add_validations,
    _apply_base_styling,
    _create_lists_sheet,
    _create_meta_sheet,
    _safe_sheet_title,
    _set_column_widths,
)
from app.features.words.workbook.importer import import_words_workbook

__all__ = ["InvalidWorkbookError", "build_words_workbook", "import_words_workbook"]


def _translation_entries_for_export(word: Word) -> str:
    if word.translation_items:
        return "\n".join(item.value for item in word.translation_items)
    return word.translations


def _example_entries_for_export(word: Word) -> str | None:
    if word.example_items:
        return "\n".join(item.value for item in word.example_items)
    return word.example


def _countability_for_export(word: Word) -> str | None:
    return _normalize_countability(word.countability)


def _load_export_topics(db: Session) -> list[Topic]:
    return list(
        db.scalars(
            select(Topic)
            .options(load_only(Topic.id, Topic.name))
            .where(Topic.deleted_at.is_(None))
            .order_by(Topic.name.asc())
        ).all()
    )


def _load_export_words_by_topic(db: Session, topic_ids: list[int]) -> dict[int, list[Word]]:
    if not topic_ids:
        return {}

    words = list(
        db.scalars(
            select(Word)
            .options(
                selectinload(Word.topics),
                selectinload(Word.translation_items),
                selectinload(Word.example_items),
            )
            .join(word_topics, Word.id == word_topics.c.word_id)
            .join(Topic, Topic.id == word_topics.c.topic_id)
            .where(Word.deleted_at.is_(None))
            .where(Topic.deleted_at.is_(None))
            .where(Topic.id.in_(topic_ids))
            .distinct()
            .order_by(Word.term.asc(), Word.id.asc())
        ).all()
    )

    words_by_topic: dict[int, list[Word]] = {topic_id: [] for topic_id in topic_ids}
    active_topic_ids = set(topic_ids)
    for word in words:
        for topic in word.topics:
            if topic.deleted_at is None and topic.id in active_topic_ids:
                words_by_topic[topic.id].append(word)
    return words_by_topic


def build_words_workbook(db: Session) -> tuple[str, bytes]:
    workbook = Workbook()
    default_sheet = workbook.active
    assert default_sheet is not None
    workbook.remove(default_sheet)

    _create_lists_sheet(workbook)

    topics = _load_export_topics(db)

    used_titles: set[str] = set()
    meta_mappings: list[tuple[str, int, str]] = []

    if not topics:
        ws = workbook.create_sheet("Vocabulary")
        for idx, (_, header) in enumerate(EXPORT_COLUMNS, start=1):
            ws.cell(row=1, column=idx, value=header)
        _set_column_widths(ws)
        _apply_base_styling(ws, 2)
        _add_dynamic_row_colors(ws, 2)
        _add_validations(workbook, ws, 2)
    else:
        words_by_topic = _load_export_words_by_topic(db, [topic.id for topic in topics])
        for topic in topics:
            words = words_by_topic.get(topic.id, [])
            sheet_title = _safe_sheet_title(topic.name, used_titles)
            meta_mappings.append((sheet_title, topic.id, topic.name))
            ws = workbook.create_sheet(sheet_title)

            for idx, (_, header) in enumerate(EXPORT_COLUMNS, start=1):
                ws.cell(row=1, column=idx, value=header)

            row_idx = 2
            for word in words:
                ws.cell(row=row_idx, column=1, value=word.knowledge_level)
                ws.cell(row=row_idx, column=2, value=word.term)
                ws.cell(row=row_idx, column=3, value=_translation_entries_for_export(word))
                ws.cell(row=row_idx, column=4, value=word.pattern)
                ws.cell(row=row_idx, column=5, value=_example_entries_for_export(word))
                ws.cell(row=row_idx, column=6, value=_countability_for_export(word))
                ws.cell(row=row_idx, column=7, value=word.part_of_speech)
                ws.cell(row=row_idx, column=8, value=word.past_simple)
                ws.cell(row=row_idx, column=9, value=word.past_participle)
                ws.cell(row=row_idx, column=10, value=word.notes)
                ws.cell(row=row_idx, column=11, value=word.id)
                row_idx += 1

            last_row = max(2, row_idx + EXTRA_EMPTY_ROWS - 1)
            _set_column_widths(ws)
            _apply_base_styling(ws, last_row)
            _add_dynamic_row_colors(ws, last_row)
            _add_validations(workbook, ws, last_row)

    _create_meta_sheet(workbook, meta_mappings)

    output = BytesIO()
    workbook.save(output)
    filename = f"lexora-vocabulary-{datetime.now(UTC).strftime('%Y%m%d-%H%M%S')}.xlsx"
    return filename, output.getvalue()
