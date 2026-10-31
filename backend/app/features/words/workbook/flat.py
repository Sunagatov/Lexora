"""One row per Word ID, with separate topic references for editing and new imports."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.features.languages.catalog import learning_language, translation_language
from app.features.topics.model import Topic
from app.features.words.repository import get_all_words
from app.features.words.workbook.format import EXPORT_COLUMNS, PRONUNCIATION_COLUMNS
from app.features.words.workbook.language import add_language_metadata


def build_flat_workbook(db: Session, *, template: bool = False, favorites_only: bool = False, topic_ids: list[int] | None = None) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Words"
    sheet.append([header for _, header in EXPORT_COLUMNS] + ["Topic IDs", "Topics"]
                 + [header for _, header in PRONUNCIATION_COLUMNS])
    topics = list(
        db.scalars(select(Topic).where(Topic.deleted_at.is_(None), Topic.language == learning_language(db)).order_by(Topic.name)).all()
    )
    if not template:
        for word in sorted(get_all_words(db, favorites_only=favorites_only, topic_ids=topic_ids), key=lambda word: word.id):
            sheet.append(
                [
                    word.knowledge_level,
                    word.term,
                    "\n".join(item.value for item in word.translation_items if (item.language or "ru") == translation_language(db)),
                    word.pattern,
                    "\n\n".join(item.value for item in word.example_items),
                    word.countability,
                    word.part_of_speech.name if word.part_of_speech else None,
                    (word.explanation_groups or {}).get("en", {}).get("definition") if translation_language(db) == "en" else word.definition,
                    word.cefr_level,
                    word.register,
                    (word.explanation_groups or {}).get("en", {}).get("notes") if translation_language(db) == "en" else word.notes,
                    word.id,
                    "\n".join(str(topic.id) for topic in word.topics if topic.deleted_at is None),
                    None,
                    *(getattr(word, name) for name, _ in PRONUNCIATION_COLUMNS),
                ]
            )
    topic_sheet = book.create_sheet("Topics")
    topic_sheet.append(["Topic ID", "Topic", "Parent ID", "Description"])
    for topic in topics:
        topic_sheet.append([topic.id, topic.name, topic.parent_topic_id, topic.description])
    for ws in book:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="37335C")
            ws.column_dimensions[cell.column_letter].width = 28
        # Vocabulary is text, including a term beginning with '='. Never emit formulas.
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    result = BytesIO()
    add_language_metadata(book, db)
    book.save(result)
    return result.getvalue()
