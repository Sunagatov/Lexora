"""Parse both topic-based exports and a flat Words/Topics workbook, without writes."""

import re
from hashlib import sha256
from io import BytesIO

from openpyxl import load_workbook

from app.features.words.constants import CEFR_LEVELS, REGISTER_VALUES
from app.features.words.workbook.cells import (
    _build_header_map,
    _ordered_unique,
    _parse_optional_int,
    _split_examples_cell,
    _split_translation_cell,
    _validate_countability,
)
from app.features.words.workbook.format import (
    HEADER_ALIASES,
    LISTS_SHEET_NAME,
    MAX_WORKBOOK_BYTES,
    META_SHEET_NAME,
    PRONUNCIATION_COLUMNS,
    InvalidWorkbookError,
)
from app.features.words.workbook.import_rules import _validate_part_of_speech
from app.features.words.workbook.language import LANGUAGE_SHEET
from app.features.words.workbook.sync_rows import POS_ALIASES

MAX_IMPORT_ROWS = 50_000


def issue(sheet: str, row: int, message: str) -> dict:
    return {"sheet": sheet, "row": row, "message": message}


def split_refs(value: str | None) -> list[str]:
    return _ordered_unique(re.split(r"[;\n]+", value or ""))


def read_import_source(content: bytes) -> dict:
    if len(content) > MAX_WORKBOOK_BYTES:
        raise InvalidWorkbookError("Workbook is too large; maximum size is 10 MB")
    try:
        book = load_workbook(BytesIO(content), read_only=True, keep_links=False)
    except Exception as exc:
        raise InvalidWorkbookError("Failed to read XLSX workbook") from exc
    source = {
        "sha256": sha256(content).hexdigest(),
        "rows": [],
        "topics": [],
        "errors": [],
        "warnings": [],
    }
    try:
        if LANGUAGE_SHEET in book:
            source["language_context"] = {str(row[0]): row[1] for row in book[LANGUAGE_SHEET].iter_rows(values_only=True) if len(row) == 2}
        if sum(ws.max_row or 0 for ws in book if ws.title != LISTS_SHEET_NAME) > MAX_IMPORT_ROWS:
            raise InvalidWorkbookError(
                "Workbook exceeds the 50,000-row limit; remove unused formatted rows"
            )
        meta = {}
        if META_SHEET_NAME in book:
            for values in book[META_SHEET_NAME].iter_rows(min_row=2, values_only=True):
                if values and values[0]:
                    name = str(values[0])
                    if len(values) < 2 or values[1] is None:
                        source["errors"].append(
                            issue(META_SHEET_NAME, 0, f"Missing topic ID for {name}")
                        )
                        continue
                    if name in meta:
                        source["errors"].append(
                            issue(META_SHEET_NAME, 0, f"Duplicate topic metadata: {name}")
                        )
                    meta[name] = values[1]
        for ws in book:
            if ws.sheet_state != "visible" or ws.title in {LISTS_SHEET_NAME, META_SHEET_NAME, LANGUAGE_SHEET}:
                continue
            if ws.max_column > 50:
                source["errors"].append(
                    issue(ws.title, 1, "Too many columns; remove unused formatted columns")
                )
                continue
            headers = _build_header_map(ws)
            seen_headers = set()
            for cell in next(ws.iter_rows(max_row=1)):
                text = " ".join(str(cell.value or "").strip().lower().split())
                canonical = next(
                    (name for name, aliases in HEADER_ALIASES.items() if text in aliases), None
                )
                if canonical in seen_headers:
                    source["errors"].append(issue(ws.title, 1, f"Duplicate column: {text}"))
                if canonical:
                    seen_headers.add(canonical)
            if ws.title == "Topics" and "term" not in headers:
                _read_topics(ws, source)
                continue
            if "term" not in headers:
                source["errors"].append(issue(ws.title, 1, "Missing Word column"))
                continue
            # Read-only worksheets must be streamed; random cell access rereads the XML.
            for row_number, cells in enumerate(
                ws.iter_rows(min_row=2, max_col=min(ws.max_column, 50)), 2
            ):
                values = {
                    name: cells[index - 1].value
                    for name, index in headers.items()
                    if index <= len(cells)
                }
                if not any(value is not None and str(value).strip() for value in values.values()):
                    continue
                try:
                    if any(cell.data_type == "f" for cell in cells):
                        raise InvalidWorkbookError(
                            "Formula cells are not accepted; paste their values instead"
                        )
                    row = _parse_row(values)
                    row.update(sheet=ws.title, row=row_number)
                    if "topic_ids" in headers or "topics" in headers:
                        row["topic_refs"] = [
                            "id:" + value for value in split_refs(_text(values.get("topic_ids")))
                        ]
                        row["topic_refs"] += [
                            "name:" + value for value in split_refs(_text(values.get("topics")))
                        ]
                    else:
                        row["topic_refs"] = (
                            ["id:" + str(meta[ws.title])]
                            if ws.title in meta
                            else ["name:" + ws.title]
                        )
                    if not row["topic_refs"]:
                        raise InvalidWorkbookError(
                            "At least one topic is required; fill Topic IDs or Topics"
                        )
                    source["rows"].append(row)
                except (InvalidWorkbookError, ValueError) as exc:
                    source["errors"].append(issue(ws.title, row_number, str(exc)))
        if not source["rows"] and not source["errors"]:
            source["errors"].append(issue("", 0, "No vocabulary rows found"))
        return source
    finally:
        book.close()


def _text(value) -> str | None:
    return str(value).strip() or None if value is not None else None


def _integer(value, label: str) -> int | None:
    result = _parse_optional_int(value, label, "Workbook", 0)
    if result is not None and result < 1:
        raise InvalidWorkbookError(f"{label} must be a positive integer")
    return result


def _parse_row(values: dict) -> dict:
    term = _text(values.get("term"))
    if not term or len(term) > 255:
        raise InvalidWorkbookError("Word must contain 1 to 255 characters")
    fields = {
        name: value
        for name in ("definition", "pattern", "notes", "cefr_level", "register")
        if (value := _text(values.get(name))) is not None
    }
    limits = {"pronunciation_ipa": 255, "pronunciation_audio_url": 512}
    for name, _ in PRONUNCIATION_COLUMNS:
        if value := _text(values.get(name)):
            if len(value) > limits.get(name, 4000):
                raise InvalidWorkbookError(f"{name} is too long")
            fields[name] = value
    if "cefr_level" in fields:
        fields["cefr_level"] = fields["cefr_level"].upper()
    for name, allowed in (("cefr_level", CEFR_LEVELS), ("register", REGISTER_VALUES)):
        if name in fields and fields[name] not in allowed:
            raise InvalidWorkbookError(f"Invalid {name}: {fields[name]}")
    pos = _text(values.get("part_of_speech"))
    if pos:
        fields["part_of_speech"] = _validate_part_of_speech(
            POS_ALIASES.get(pos.lower(), pos.lower()), "Workbook", 0
        )
    if countability := _text(values.get("countability")):
        fields["countability"] = _validate_countability(countability, "Workbook", 0)
    if translations := _text(values.get("translations")):
        fields["translations"] = _ordered_unique(_split_translation_cell(translations))
    if examples := _text(values.get("examples")):
        fields["examples"] = _split_examples_cell(examples)
    knowledge = _integer(values.get("knowledge_level"), "Knowledge")
    if knowledge is not None and knowledge > 5:
        raise InvalidWorkbookError("Knowledge must be between 1 and 5")
    return {"term": term, "word_id": _integer(values.get("word_id"), "Word ID"), "fields": fields}


def _read_topics(ws, source: dict) -> None:
    header = {
        str(cell.value).strip().lower(): i for i, cell in enumerate(next(ws.iter_rows(max_row=1)))
    }
    if "topic" not in header:
        source["errors"].append(issue(ws.title, 1, "Topics sheet requires a Topic column"))
        return
    for number, raw_cells in enumerate(ws.iter_rows(min_row=2), 2):
        cells = [cell.value for cell in raw_cells]
        if not any(value is not None and str(value).strip() for value in cells):
            continue
        values = {
            key: cells[index] if index < len(cells) else None for key, index in header.items()
        }
        try:
            if any(cell.data_type == "f" for cell in raw_cells):
                raise InvalidWorkbookError("Formula cells are not accepted in topic definitions")
            name = _text(values.get("topic"))
            if not name or len(name) > 200:
                raise InvalidWorkbookError("Topic must contain 1 to 200 characters")
            source["topics"].append(
                {
                    "name": name,
                    "id": _integer(values.get("topic id"), "Topic ID"),
                    "parent_topic_id": _integer(values.get("parent id"), "Parent ID"),
                    "description": _text(values.get("description")),
                    "sheet": ws.title,
                    "row": number,
                }
            )
        except InvalidWorkbookError as exc:
            source["errors"].append(issue(ws.title, number, str(exc)))
