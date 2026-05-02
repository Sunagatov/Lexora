from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.features.topics.model import Topic
from app.features.words.api import existing_normalized_terms, sync_word_multivalue_fields
from app.features.words.domain import assert_no_duplicate_word
from app.features.words.model import Word
from app.features.words.progress import record_level_change
from app.features.words.schemas import WorkbookImportResponse, WorkbookImportSheetSummary
from app.features.words.workbook.cells import _build_header_map, _read_meta_topic_refs
from app.features.words.workbook.format import (
    LISTS_SHEET_NAME,
    MAX_WORKBOOK_BYTES,
    META_SHEET_NAME,
    InvalidWorkbookError,
)
from app.features.words.workbook.import_apply import (
    WorkbookImportOps,
    WorkbookImportFlags,
    WorkbookSheetContext,
    create_new_word_from_row,
    update_existing_word_from_row,
)
from app.features.words.workbook.import_rows import _read_row_data
from app.features.words.workbook.import_rules import (
    _find_existing_word,
    _get_topic,
    _should_validate_existing_word_duplicate,
)


def _import_sheet(
    db: Session,
    ws,
    header_map: dict[str, int],
    topic: Topic,
    seen_existing_words: dict[int, tuple[str, tuple[object, ...]]] | None = None,
) -> WorkbookImportSheetSummary:
    if seen_existing_words is None:
        seen_existing_words = {}

    enabled_header_map = _enabled_header_map(header_map)
    context = WorkbookSheetContext(
        flags=WorkbookImportFlags(
            has_knowledge="knowledge_level" in header_map,
            has_pattern="pattern" in header_map,
            has_examples="examples" in header_map,
            has_countability="countability" in header_map,
            has_part_of_speech="part_of_speech" in header_map,
            has_definition="definition" in header_map,
            has_cefr_level="cefr_level" in header_map,
            has_register="register" in header_map,
            has_notes="notes" in header_map,
        ),
        sheet_name=ws.title,
        seen_existing_words=seen_existing_words,
        ops=WorkbookImportOps(
            should_validate_existing_word_duplicate=_should_validate_existing_word_duplicate,
            assert_no_duplicate_word=assert_no_duplicate_word,
            existing_normalized_terms=existing_normalized_terms,
            record_level_change=record_level_change,
            sync_word_multivalue_fields=sync_word_multivalue_fields,
            word_cls=Word,
        ),
    )

    created = 0
    updated = 0
    skipped = 0
    seen_word_ids: set[int] = set()

    for row_idx in range(2, ws.max_row + 1):
        row = _read_row_data(ws, row_idx, enabled_header_map)
        if row is None:
            continue

        _track_seen_word_id(seen_word_ids, row.word_id, ws.title, row_idx)
        existing = _find_existing_word(db, topic.id, row.word_id, row.term)

        if existing is not None:
            update_existing_word_from_row(
                db,
                existing=existing,
                row=row,
                topic=topic,
                row_idx=row_idx,
                context=context,
            )
            updated += 1
            continue

        create_new_word_from_row(
            db,
            row=row,
            topic=topic,
            row_idx=row_idx,
            context=context,
        )
        created += 1

    return WorkbookImportSheetSummary(
        sheet_name=ws.title,
        topic_name=topic.name,
        created=created,
        updated=updated,
        skipped=skipped,
    )


def import_words_workbook(db: Session, content: bytes) -> WorkbookImportResponse:
    if len(content) > MAX_WORKBOOK_BYTES:
        raise InvalidWorkbookError(
            f"Workbook is too large; maximum size is {MAX_WORKBOOK_BYTES // (1024 * 1024)} MB"
        )

    try:
        workbook = load_workbook(filename=BytesIO(content))
    except Exception as exc:
        raise InvalidWorkbookError("Failed to read XLSX workbook") from exc

    meta_topic_refs = _read_meta_topic_refs(workbook)

    summaries: list[WorkbookImportSheetSummary] = []
    total_created = 0
    total_updated = 0
    total_skipped = 0
    seen_existing_words: dict[int, tuple[str, tuple[object, ...]]] = {}

    for ws in workbook.worksheets:
        if ws.sheet_state != "visible":
            continue
        if ws.title in {LISTS_SHEET_NAME, META_SHEET_NAME}:
            continue

        header_map = _build_header_map(ws)
        if "term" not in header_map or "translations" not in header_map:
            continue

        topic_id, topic_name = meta_topic_refs.get(ws.title, (None, None))
        resolved_topic_name = topic_name or ws.title
        topic = _get_topic(db, resolved_topic_name, topic_id=topic_id)

        summary = _import_sheet(db, ws, header_map, topic, seen_existing_words)
        summaries.append(summary)
        total_created += summary.created
        total_updated += summary.updated
        total_skipped += summary.skipped

    if not summaries:
        raise InvalidWorkbookError("No importable sheets found in workbook")

    db.commit()
    return WorkbookImportResponse(
        created=total_created,
        updated=total_updated,
        skipped=total_skipped,
        sheets=summaries,
    )


def _enabled_header_map(header_map: dict[str, int]) -> dict[str, int]:
    optional_headers = {
        "word_id",
        "knowledge_level",
        "pattern",
        "examples",
        "countability",
        "part_of_speech",
        "definition",
        "cefr_level",
        "register",
        "notes",
    }
    return {
        key: value
        for key, value in header_map.items()
        if key in {"term", "translations"} or key in optional_headers
    }


def _track_seen_word_id(
    seen_word_ids: set[int],
    word_id: int | None,
    sheet_name: str,
    row_idx: int,
) -> None:
    if word_id is None:
        return
    if word_id in seen_word_ids:
        raise InvalidWorkbookError(f"{sheet_name}, row {row_idx}: duplicate Word ID {word_id} in this sheet")
    seen_word_ids.add(word_id)
