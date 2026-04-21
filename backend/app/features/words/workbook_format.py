import re
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

LISTS_SHEET_NAME = "Lists"
META_SHEET_NAME = "__lexora_meta"
MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
EXTRA_EMPTY_ROWS = 30

EXPORT_COLUMNS = [
    ("knowledge_level", "Knowledge"),
    ("term", "Word"),
    ("translations", "Russian translations"),
    ("pattern", "Typical prepositions / patterns"),
    ("examples", "Examples (EN + RU)"),
    ("countability", "Countability"),
    ("part_of_speech", "Part of speech"),
    ("past_simple", "Past simple"),
    ("past_participle", "Past participle"),
    ("notes", "Notes"),
    ("word_id", "Word ID"),
]

HEADER_ALIASES: dict[str, set[str]] = {
    "knowledge_level": {"knowledge", "level"},
    "term": {"word", "term", "base form", "phrase"},
    "translations": {
        "russian translations",
        "translations",
        "translation",
        "russian translation",
    },
    "pattern": {
        "typical prepositions / patterns",
        "pattern",
        "patterns",
        "prepositions / patterns",
        "typical position / usage",
    },
    "examples": {"examples (en + ru)", "examples", "example", "example (en+ru)"},
    "countability": {"countability"},
    "part_of_speech": {"part of speech", "pos"},
    "past_simple": {"past simple"},
    "past_participle": {"past participle"},
    "notes": {"notes", "note", "meaning / usage note"},
    "word_id": {"word id", "id"},
}

PART_OF_SPEECH_VALUES = [
    "noun",
    "verb",
    "adjective",
    "adverb",
    "phrase",
    "preposition",
    "other",
]
COUNTABILITY_VALUES = ["Countable", "Uncountable", "Both", "Plural", "Collective"]
COUNTABILITY_ALIASES = {
    "countable": "Countable",
    "uncountable": "Uncountable",
    "both": "Both",
    "plural": "Plural",
    "collective": "Collective",
}

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)
LEVEL_FILLS = {
    1: PatternFill(fill_type="solid", fgColor="FEE2E2"),
    2: PatternFill(fill_type="solid", fgColor="FEF3C7"),
    3: PatternFill(fill_type="solid", fgColor="DBEAFE"),
    4: PatternFill(fill_type="solid", fgColor="DCFCE7"),
    5: PatternFill(fill_type="solid", fgColor="E5E7EB"),
}


class InvalidWorkbookError(Exception):
    pass


def _safe_sheet_title(raw_title: str, used_titles: set[str]) -> str:
    cleaned = re.sub(r"[:\\/?*\[\]]", " ", raw_title).strip()
    if not cleaned:
        cleaned = "Topic"
    cleaned = re.sub(r"\s+", " ", cleaned)

    base = cleaned[:31]
    candidate = base
    suffix = 2
    while candidate in used_titles:
        suffix_text = f" ({suffix})"
        candidate = f"{base[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1

    used_titles.add(candidate)
    return candidate


def _set_column_widths(ws) -> None:
    widths = {
        "A": 12,
        "B": 28,
        "C": 36,
        "D": 28,
        "E": 42,
        "F": 16,
        "G": 18,
        "H": 18,
        "I": 20,
        "J": 28,
        "K": 12,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.column_dimensions["K"].hidden = True


def _apply_base_styling(ws, last_row: int) -> None:
    for row in ws.iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=len(EXPORT_COLUMNS)):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{last_row}"


def _add_dynamic_row_colors(ws, last_row: int) -> None:
    if last_row < 2:
        return

    for level, fill in LEVEL_FILLS.items():
        ws.conditional_formatting.add(
            f"A2:J{last_row}",
            FormulaRule(formula=[f"$A2={level}"], stopIfTrue=False, fill=fill),
        )


def _add_validations(workbook: Workbook, ws, last_row: int) -> None:
    knowledge_dv = DataValidation(
        type="list", formula1=f"={LISTS_SHEET_NAME}!$A$1:$A$5", allow_blank=True
    )
    countability_dv = DataValidation(
        type="list",
        formula1=f"={LISTS_SHEET_NAME}!$B$1:$B${len(COUNTABILITY_VALUES)}",
        allow_blank=True,
    )
    pos_dv = DataValidation(
        type="list", formula1=f"={LISTS_SHEET_NAME}!$C$1:$C$7", allow_blank=True
    )

    ws.add_data_validation(knowledge_dv)
    ws.add_data_validation(countability_dv)
    ws.add_data_validation(pos_dv)

    knowledge_dv.add(f"A2:A{last_row}")
    countability_dv.add(f"F2:F{last_row}")
    pos_dv.add(f"G2:G{last_row}")


def _create_lists_sheet(workbook: Workbook) -> None:
    ws = workbook.create_sheet(LISTS_SHEET_NAME)
    for idx, value in enumerate(["1", "2", "3", "4", "5"], start=1):
        ws.cell(row=idx, column=1, value=value)
    for idx, value in enumerate(COUNTABILITY_VALUES, start=1):
        ws.cell(row=idx, column=2, value=value)
    for idx, value in enumerate(PART_OF_SPEECH_VALUES, start=1):
        ws.cell(row=idx, column=3, value=value)
    ws.sheet_state = "hidden"


def _create_meta_sheet(workbook: Workbook, mappings: list[tuple[str, int, str]]) -> None:
    ws = workbook.create_sheet(META_SHEET_NAME)
    ws.append(["sheet_name", "topic_id", "topic_name", "exported_at"])
    exported_at = datetime.now(UTC).isoformat()
    for sheet_name, topic_id, topic_name in mappings:
        ws.append([sheet_name, topic_id, topic_name, exported_at])
    ws.sheet_state = "hidden"
