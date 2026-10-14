import json
from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session

from app.features.smart_review.model import StudyQueue, StudyQueueItem
from app.features.topics.model import Topic
from app.features.words.model import PartsOfSpeech, Word, WordExample, WordTranslation, WordVerbForm
from app.features.words.progress import WordProgressEvent
from app.features.words.workbook.cells import _split_examples_cell
from app.features.words.workbook.format import InvalidWorkbookError
from app.features.words.workbook.importer import import_words_workbook
from app.features.words.workbook.service import build_words_workbook
from app.features.words.workbook.sync import (
    apply_sync_plan,
    build_restore_plan,
    build_sync_plan,
    content_state,
    protected_digest,
    write_recovery_snapshot,
)
from app.features.words.workbook.sync_rows import positive_id, read_sync_workbook
from app.shared.db import Base

HEADERS = [
    "Knowledge",
    "Word",
    "Russian translations",
    "Typical prepositions / patterns",
    "Examples (EN + RU)",
    "Countability",
    "Part of speech",
    "Definition",
    "CEFR",
    "Register",
    "Notes",
    "Word ID",
]


def workbook_bytes(rows, extra_sheets=None):
    book = Workbook()
    book.remove(book.active)
    meta = book.create_sheet("__lexora_meta")
    meta.append(["sheet_name", "topic_id", "topic_name", "exported_at"])
    for name, topic_id, entries in [("Animals", 1, rows), *(extra_sheets or [])]:
        ws = book.create_sheet(name)
        ws.append(HEADERS)
        for row in entries:
            ws.append([row.get(header) for header in HEADERS])
        meta.append([name, topic_id, name, "2026-10-14"])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def row(**overrides):
    return {
        "Knowledge": 1,
        "Word": "bank",
        "Word ID": 10,
        "Russian translations": "банк\nберег\nбанк",
        "Part of speech": "noun",
        "Definition": "A financial institution.",
        "CEFR": "A2",
        "Register": "neutral",
        "Countability": "countable",
        "Examples (EN + RU)": "EN: I went to the bank.\nRU: Я пошёл в банк.",
        **overrides,
    }


@pytest.fixture
def database():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        noun, verb = PartsOfSpeech(id=1, name="noun"), PartsOfSpeech(id=2, name="verb")
        topic = Topic(id=1, name="Animals", slug="animals")
        other = Topic(id=2, name="Finance", slug="finance", parent_topic_id=1)
        word = Word(
            id=10,
            term="bank",
            language="en",
            topics=[topic],
            part_of_speech=noun,
            knowledge_level=4,
            definition="Old definition",
            pattern="at the bank",
            notes="Keep my personal note",
            image_url="https://example.test/image.png",
            pronunciation_ipa="bæŋk",
            countability="countable",
            cefr_level="B1",
            register="formal",
            verb_form=WordVerbForm(past_simple="preserved"),
        )
        word.translation_items = [WordTranslation(position=0, value="старый перевод")]
        word.example_items = [
            WordExample(position=0, value="I went to the bank."),
            WordExample(position=1, value="An old valid example."),
        ]
        db.add_all(
            [
                noun,
                verb,
                PartsOfSpeech(id=3, name="phrase"),
                PartsOfSpeech(id=4, name="other"),
                topic,
                other,
                word,
            ]
        )
        now = datetime.now(UTC)
        queue = StudyQueue(
            id=1,
            expires_at=now + timedelta(days=1),
            total_count=1,
            completed_count=1,
            items=[StudyQueueItem(word=word, position=0, is_completed=True, completed_at=now)],
        )
        db.add_all([queue, WordProgressEvent(word=word, old_level=1, new_level=4, source="manual")])
        db.commit()
        yield db
    engine.dispose()


def test_content_sync_is_dry_run_idempotent_and_preserves_progress(database, tmp_path):
    source = read_sync_workbook(
        workbook_bytes(
            [row(Notes="Topic: Animals")],
            [("Finance", 2, [row(Knowledge=99, Notes="Topic: Finance")])],
        )
    )
    word = database.get(Word, 10)
    before = content_state(word)
    protected = protected_digest(database)
    plan = build_sync_plan(database, source)
    assert not plan["errors"] and not plan["conflicts"]
    assert plan["summary"]["unique_words"] == 1
    assert plan["summary"]["created"] == 0
    assert content_state(word) == before
    assert not database.dirty
    backup = tmp_path / "before.json"
    write_recovery_snapshot(backup, plan)
    result = apply_sync_plan(database, plan)
    database.commit()
    assert result["progress_preserved"]
    assert protected_digest(database) == protected
    assert word.knowledge_level == 4
    assert [t.value for t in word.translation_items] == ["банк", "берег"]
    assert word.example_items[0].value == "EN: I went to the bank.\nRU: Я пошёл в банк."
    assert [e.value for e in word.example_items][1:] == ["EN: An old valid example."]
    assert word.cefr_level == "A2" and word.register == "neutral"
    assert {t.id for t in word.topics} == {1, 2}
    assert word.notes == "Keep my personal note\nTopic: Animals\nTopic: Finance"
    again = build_sync_plan(database, source)
    assert again["changes"] == []
    assert again["summary"]["unchanged"] == 1
    restore = build_restore_plan(database, backup)
    apply_sync_plan(database, restore)
    database.commit()
    assert content_state(word) == before
    assert protected_digest(database) == protected
    assert build_restore_plan(database, backup)["changes"] == []


def test_conflicting_shared_ids_are_reported_and_block_all_writes(database):
    source = read_sync_workbook(
        workbook_bytes([row()], [("Finance", 2, [row(Definition="Different sense.")])])
    )
    plan = build_sync_plan(database, source)
    assert source.conflicts[0]["field"] == "definition"
    assert len(source.conflicts[0]["variants"]) == 2
    before = content_state(database.get(Word, 10))
    with pytest.raises(InvalidWorkbookError, match="Unresolved"):
        apply_sync_plan(database, plan)
    assert content_state(database.get(Word, 10)) == before
    assert plan["summary"]["skipped"] == 1


def test_explicit_resolution_chooses_only_an_observed_variant(database):
    book = workbook_bytes([row()], [("Finance", 2, [row(Definition="Different sense.")])])
    source = read_sync_workbook(book, {"10": {"definition": "Different sense."}})
    assert not source.conflicts
    plan = build_sync_plan(database, source)
    apply_sync_plan(database, plan)
    assert database.get(Word, 10).definition == "Different sense."
    invalid = read_sync_workbook(book, {"10": {"definition": "Not in workbook"}})
    assert invalid.errors


def test_blank_cells_preserve_useful_content_and_metadata(database):
    source = read_sync_workbook(workbook_bytes([{"Word": "bank", "Word ID": 10, "Knowledge": 2}]))
    plan = build_sync_plan(database, source)
    assert not plan["changes"]
    assert database.get(Word, 10).pattern == "at the bank"
    assert database.get(Word, 10).knowledge_level == 4


def test_pos_alias_keeps_precise_label_in_notes_and_updates_relationship(database):
    source = read_sync_workbook(workbook_bytes([row(**{"Part of speech": "proper noun"})]))
    plan = build_sync_plan(database, source)
    apply_sync_plan(database, plan)
    assert "Workbook part of speech: proper noun" in database.get(Word, 10).notes
    assert source.normalizations[0]["to"] == "noun"


@pytest.mark.parametrize("value", [None, True, 0, -1, "1.5", 1.5, "oops"])
def test_rejects_malformed_ids(value):
    with pytest.raises(ValueError):
        positive_id(value)


@pytest.mark.parametrize(
    "changed",
    [
        {"Word ID": 999},
        {"Part of speech": "made up"},
        {"CEFR": "C9"},
        {"Register": "invented"},
        {"Countability": "unknown"},
        {"Definition": "=1+1"},
        {"Examples (EN + RU)": "RU: Без английского предложения."},
    ],
)
def test_validation_never_allows_malformed_content_to_be_applied(database, changed):
    plan = build_sync_plan(database, read_sync_workbook(workbook_bytes([row(**changed)])))
    assert plan["errors"]
    before = content_state(database.get(Word, 10))
    with pytest.raises(InvalidWorkbookError):
        apply_sync_plan(database, plan, skip_conflicts=True)
    assert content_state(database.get(Word, 10)) == before


def test_unknown_and_deleted_topic_references_are_rejected(database):
    source = read_sync_workbook(workbook_bytes([row()], [("Unknown", 999, [row()])]))
    plan = build_sync_plan(database, source)
    assert plan["errors"][0]["topic_ids"] == [999]
    database.get(Topic, 1).deleted_at = datetime.now(UTC)
    database.flush()
    plan = build_sync_plan(database, read_sync_workbook(workbook_bytes([row()])))
    assert plan["errors"]


def test_updated_pos_is_persisted_by_regular_import_along_with_cefr_and_register(database):
    content = workbook_bytes(
        [row(**{"Word": "bank", "Part of speech": "verb", "Countability": None})]
    )
    import_words_workbook(database, content)
    database.expire_all()
    word = database.get(Word, 10)
    assert word.part_of_speech.name == "verb"
    assert word.part_of_speech_id == 2
    assert word.cefr_level == "A2"
    assert word.register == "neutral"
    assert len(word.example_items) == 1


def test_different_parts_of_speech_keep_distinct_word_ids(database):
    other = Word(
        id=11,
        term="bank",
        part_of_speech=database.get(PartsOfSpeech, 2),
        topics=[database.get(Topic, 1)],
        knowledge_level=3,
    )
    database.add(other)
    database.commit()
    plan = build_sync_plan(
        database,
        read_sync_workbook(
            workbook_bytes([row(**{"Part of speech": "verb", "Countability": None})])
        ),
    )
    assert any("collide" in error["error"] for error in plan["conflicts"])
    assert database.get(Word, 11).knowledge_level == 3


def test_stale_plan_and_corrupted_backup_are_rejected(database, tmp_path):
    plan = build_sync_plan(database, read_sync_workbook(workbook_bytes([row()])))
    backup = tmp_path / "backup.json"
    write_recovery_snapshot(backup, plan)
    database.get(Word, 10).definition = "Later edit"
    database.flush()
    with pytest.raises(InvalidWorkbookError, match="changed after planning"):
        apply_sync_plan(database, plan)
    payload = json.loads(backup.read_text())
    payload["payload"]["changes"][0]["word_id"] = 999
    backup.write_text(json.dumps(payload))
    with pytest.raises(InvalidWorkbookError, match="corrupted"):
        build_restore_plan(database, backup)


def test_restore_refuses_to_overwrite_later_content_edits(database, tmp_path):
    plan = build_sync_plan(database, read_sync_workbook(workbook_bytes([row()])))
    backup = tmp_path / "backup.json"
    write_recovery_snapshot(backup, plan)
    apply_sync_plan(database, plan)
    database.commit()
    database.get(Word, 10).definition = "Newer manual content"
    database.commit()
    restore = build_restore_plan(database, backup)
    assert restore["errors"]
    with pytest.raises(InvalidWorkbookError):
        apply_sync_plan(database, restore)


def test_export_roundtrip_retains_paired_examples_and_progress(database):
    source = read_sync_workbook(workbook_bytes([row()]))
    apply_sync_plan(database, build_sync_plan(database, source))
    database.commit()
    _, content = build_words_workbook(database)
    exported = load_workbook(BytesIO(content))
    assert exported["Animals"].cell(2, 1).value == 4
    reread = read_sync_workbook(content)
    assert not reread.errors and not reread.conflicts
    assert build_sync_plan(database, reread)["changes"] == []


def test_pair_parser_rejects_orphan_ru_and_preserves_english_only_examples():
    assert _split_examples_cell("One.\nTwo.") == ["One.", "Two."]
    assert _split_examples_cell("EN: One.\nRU: Один.\nEN: Two.\nRU: Два.") == [
        "EN: One.\nRU: Один.",
        "EN: Two.\nRU: Два.",
    ]
    with pytest.raises(InvalidWorkbookError):
        _split_examples_cell("RU: Один.")


def test_skipping_conflicts_leaves_entire_conflicted_word_unchanged(database):
    other = Word(
        id=11,
        term="charge",
        part_of_speech=database.get(PartsOfSpeech, 1),
        topics=[database.get(Topic, 1)],
        knowledge_level=3,
    )
    database.add(other)
    database.commit()
    content = workbook_bytes(
        [row(), row(**{"Word ID": 11, "Word": "charge"})],
        [("Finance", 2, [row(Definition="Conflicting meaning")])],
    )
    source = read_sync_workbook(content)
    plan = build_sync_plan(database, source)
    before = content_state(database.get(Word, 10))
    apply_sync_plan(database, plan, skip_conflicts=True)
    database.commit()
    assert content_state(database.get(Word, 10)) == before
    assert database.get(Word, 11).definition == "A financial institution."
    assert database.get(Word, 11).knowledge_level == 3
    assert build_sync_plan(database, source)["changes"] == []


def test_related_notes_only_propose_exact_existing_topics_until_included(database):
    source = read_sync_workbook(workbook_bytes([row(Notes="Related: Finance, Imaginary")]))
    plan = build_sync_plan(database, source)
    assert plan["topic_suggestions"] == [{"word_id": 10, "topic_id": 2, "topic_name": "Finance"}]
    assert plan["unmatched_topic_labels"] == [{"word_id": 10, "label": "Imaginary"}]
    assert "topic_ids" not in plan["changes"][0]["fields"]
    included = build_sync_plan(database, source, include_related_topics=True)
    protected = protected_digest(database)
    apply_sync_plan(database, included)
    database.commit()
    assert {topic.id for topic in database.get(Word, 10).topics} == {1, 2}
    assert protected_digest(database) == protected


def test_database_failure_rolls_back_the_complete_batch(database):
    source = read_sync_workbook(workbook_bytes([row()]))
    plan = build_sync_plan(database, source)
    word = database.get(Word, 10)
    before = content_state(word)
    protected = protected_digest(database)
    connection = database.get_bind()

    def fail_insert(_connection, _cursor, statement, _parameters, _context, _many):
        if statement.startswith("INSERT INTO word_translations"):
            raise RuntimeError("Simulated write failure")

    event.listen(connection, "before_cursor_execute", fail_insert)
    try:
        with pytest.raises(RuntimeError, match="Simulated"):
            apply_sync_plan(database, plan)
        database.rollback()
    finally:
        event.remove(connection, "before_cursor_execute", fail_insert)
    assert content_state(database.get(Word, 10)) == before
    assert protected_digest(database) == protected
