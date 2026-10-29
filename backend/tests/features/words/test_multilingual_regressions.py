from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.features.languages.catalog import configure_language
from app.features.topics.model import Topic
from app.features.words.model import PartsOfSpeech, Word, WordExample, WordTranslation
from app.features.words.progress import WordProgressEvent
from app.features.words.repository import get_all_words, update_word
from app.features.words.schemas import WordUpdate
from app.features.words.workbook.flat import build_flat_workbook
from app.features.words.workbook.import_batches import (
    StaleImportError,
    apply_import,
    preview_import,
    recovery_record,
)
from app.features.words.workbook.import_plan import build_import_plan
from app.features.words.workbook.import_source import read_import_source
from app.features.words.workbook.service import build_words_workbook
from app.shared.db import Base


@pytest.fixture
def library():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        configure_language(db, "zh", "ru")
        word = Word(term="你好", language="zh", knowledge_level=4, definition="Приветствие",
                    part_of_speech=PartsOfSpeech(name="phrase"),
                    topics=[Topic(name="Greetings", slug="greetings", language="zh")],
                    transliteration="nǐ hǎo", pronunciation_ipa="ni˨˩˦ xau˨˩˦",
                    pronunciation_easy="ни³ хао³", pronunciation_note_ru="Тоны важны",
                    pronunciation_audio_url="https://example.test/hello.mp3",
                    translation_items=[WordTranslation(language="ru", position=0, value="привет")],
                    example_items=[WordExample(position=i, value=f"ZH: 你好{i}\nRU: Привет{i}\nEN: Hello{i}") for i in range(3)])
        db.add(word)
        db.commit()
        yield db, word
    engine.dispose()


def test_rename_clears_only_stale_pronunciation_and_keeps_new_readings(library):
    db, word = library
    update_word(db, word, WordUpdate(term="谢谢", transliteration="xièxie", pronunciation_ipa="new IPA"))
    assert word.transliteration == "xièxie" and word.pronunciation_ipa == "new IPA"
    assert word.pronunciation_audio_url is word.pronunciation_easy is word.pronunciation_note_ru is None
    assert word.knowledge_level == 4
    update_word(db, word, WordUpdate(definition="Другая формулировка"))
    assert word.transliteration == "xièxie"


def test_rename_drops_exported_old_readings_even_if_resubmitted(library):
    db, word = library
    update_word(db, word, WordUpdate(term="谢谢", transliteration="nǐ hǎo", pronunciation_ipa=word.pronunciation_ipa))
    assert word.transliteration is word.pronunciation_ipa is word.pronunciation_audio_url is None


@pytest.mark.parametrize("search", ["ni hao", "nihao", "nǐ hǎo", "NI HAO", "ni3hao3"])
def test_mandarin_search_accepts_plain_spaced_compact_and_numbered_pinyin(library, search):
    db, word = library
    words, total = get_all_words(db, search=search, page=1, page_size=50)
    assert total == 1 and words[0].id == word.id


@pytest.mark.parametrize("search", ["lüe", "lve", "lu:e"])
def test_pinyin_umlaut_is_distinct_from_u(library, search):
    db, _ = library
    db.add(Word(term="略", language="zh", transliteration="lüè"))
    db.commit()
    words, _ = get_all_words(db, search=search, page=1, page_size=50)
    assert [w.term for w in words] == ["略"]


@pytest.mark.parametrize("flat", [True, False])
def test_both_exports_roundtrip_pronunciation_and_examples_without_updates(library, flat):
    db, _ = library
    data = build_flat_workbook(db) if flat else build_words_workbook(db)[1]
    source = read_import_source(data)
    assert not source["errors"]
    plan = build_import_plan(db, source, mode="update", minimum_translations=1)
    assert not plan["errors"] and plan["summary"]["unchanged"] == 1
    assert plan["entries"][0]["after"]["transliteration"] == "nǐ hǎo"
    assert len(plan["entries"][0]["after"]["examples"]) == 3


def test_workbook_import_updates_reading_with_receipt_and_preserves_progress(library):
    db, word = library
    book = load_workbook(BytesIO(build_flat_workbook(db)))
    headers = {cell.value: cell.column for cell in book["Words"][1]}
    book["Words"].cell(2, headers["Russian reading hint"], "новая подсказка")
    data = BytesIO(); book.save(data)
    preview = preview_import(db, read_import_source(data.getvalue()), mode="update", minimum_translations=1)
    assert not preview["errors"]
    apply_import(db, preview["batch_id"], preview["plan_hash"])
    db.refresh(word)
    assert word.pronunciation_easy == "новая подсказка" and word.knowledge_level == 4
    assert db.scalar(select(func.count()).select_from(WordProgressEvent)) == 0
    recovery = recovery_record(db, preview["batch_id"])["payload"]["changes"][0]
    assert recovery["before"]["pronunciation_easy"] == "ни³ хао³"
    assert recovery["after"]["pronunciation_easy"] == "новая подсказка"
    assert apply_import(db, preview["batch_id"], preview["plan_hash"])["replayed"]


def test_import_rejects_reading_edits_after_preview(library):
    db, word = library
    preview = preview_import(db, read_import_source(build_flat_workbook(db)), mode="update", minimum_translations=1)
    word.pronunciation_easy = "manually changed"
    db.commit()
    with pytest.raises(StaleImportError):
        apply_import(db, preview["batch_id"], preview["plan_hash"])


def test_new_workbook_rows_import_their_readings(library):
    db, _ = library
    source = read_import_source(build_flat_workbook(db))
    row = source["rows"][0]
    row.update(word_id=None, term="谢谢")
    row["fields"].update(translations=["спасибо"], definition="Благодарность", transliteration="xièxie",
                         pronunciation_easy="се⁴ се⁰", pronunciation_ipa="new", pronunciation_audio_url="https://example.test/thanks.mp3")
    preview = preview_import(db, source, mode="add", minimum_translations=1)
    assert not preview["errors"]
    result = apply_import(db, preview["batch_id"], preview["plan_hash"])
    word = db.get(Word, result["created_word_ids"][0])
    assert word.transliteration == "xièxie" and word.pronunciation_easy == "се⁴ се⁰"
    assert word.knowledge_level == 1


def test_older_workbooks_preserve_readings_and_rename_previews_clear_stale_ones(library):
    db, _word = library
    book = load_workbook(BytesIO(build_flat_workbook(db)))
    book["Words"].delete_cols(15, 5)
    book["Words"]["H2"] = "Новое определение"
    data = BytesIO(); book.save(data)
    plan = build_import_plan(db, read_import_source(data.getvalue()), mode="update", minimum_translations=1)
    assert plan["entries"][0]["after"]["transliteration"] == "nǐ hǎo"
    source = read_import_source(build_flat_workbook(db))
    source["rows"][0]["term"] = "谢谢"
    plan = build_import_plan(db, source, mode="update", minimum_translations=1)
    assert plan["entries"][0]["after"]["transliteration"] is None
    assert plan["entries"][0]["after"]["pronunciation_easy"] is None
