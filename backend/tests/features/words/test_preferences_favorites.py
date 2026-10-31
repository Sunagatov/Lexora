from datetime import UTC, datetime
from io import BytesIO
from zipfile import ZipFile

import pytest
from fastapi import Depends, FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.features.languages.catalog import configure_language, learning_language
from app.features.preferences.model import AppPreferences
from app.features.preferences.router import router as preferences_router
from app.features.preferences.schemas import Preferences, PreferencesPatch
from app.features.preferences.service import get_preferences, update_preferences
from app.features.topics.model import Topic
from app.features.words.model import Word, WordTranslation
from app.features.words.progress import WordProgressEvent
from app.features.words.repository import (
    get_all_words,
    get_ordered_word_ids,
    get_word_knowledge_summary,
    update_word,
)
from app.features.words.router import router as words_router
from app.features.words.schemas import WordUpdate
from app.shared.db import Base
from app.shared.deps import get_db, verify_csrf, verify_session


@pytest.fixture
def db():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        yield db
    engine.dispose()


def test_preferences_are_global_partial_persistent_and_validate_before_writing(db):
    assert get_preferences(db) == Preferences()
    assert db.scalar(select(func.count()).select_from(AppPreferences)) == 0
    first = update_preferences(
        db, PreferencesPatch(page_size=100, swipes=False, pronunciation_modes={"zh": "pinyin"})
    )
    configure_language(db, "es")
    second = update_preferences(
        db, PreferencesPatch(theme="light", pronunciation_modes={"es": "ipa"})
    )
    assert second.page_size == first.page_size == 100
    assert second.swipes is False
    assert second.pronunciation_modes == {"zh": "pinyin", "es": "ipa"}
    with Session(db.get_bind()) as other_device:
        assert get_preferences(other_device) == second
    for invalid in [
        {"page_size": 20},
        {"language": "xx"},
        {"swipes": None},
        {"swipes": "false"},
        {"unknown": True},
    ]:
        with pytest.raises(ValidationError):
            PreferencesPatch.model_validate(invalid)
    assert get_preferences(db) == second


def test_preferences_require_login_and_csrf(db):
    app = FastAPI()
    app.include_router(
        preferences_router, dependencies=[Depends(verify_session), Depends(verify_csrf)]
    )
    app.dependency_overrides[get_db] = lambda: db
    with TestClient(app) as client:
        assert client.get("/api/preferences").status_code == 401
        assert client.patch("/api/preferences", json={"swipes": False}).status_code in (401, 403)


def test_favorites_preserve_learning_and_are_scoped_before_sorting_pagination_and_summary(db):
    en_topic = Topic(name="English", slug="english", language="en")
    es_topic = Topic(name="Spanish", slug="spanish", language="es")
    first = Word(
        term="patient", language="en", knowledge_level=1, cefr_level="A1", topics=[en_topic]
    )
    second = Word(
        term="peaceful", language="en", knowledge_level=4, cefr_level="A2", topics=[en_topic]
    )
    foreign = Word(
        term="hola", language="es", knowledge_level=1, is_favorite=True, topics=[es_topic]
    )
    deleted = Word(
        term="gone",
        language="en",
        knowledge_level=1,
        is_favorite=True,
        deleted_at=datetime.now(UTC),
    )
    db.add_all([first, second, foreign, deleted])
    db.commit()
    configure_language(db, "en")
    before_topics = [topic.id for topic in first.topics]
    update_word(db, first, WordUpdate(is_favorite=True))
    update_word(db, second, WordUpdate(is_favorite=True))
    words, total = get_all_words(db, favorites_only=True, sort="level-asc", page=1, page_size=1)
    assert total == 2 and [word.id for word in words] == [first.id]
    assert get_ordered_word_ids(db, favorites_only=True, sort="level-asc") == [first.id, second.id]
    assert get_word_knowledge_summary(db, favorites_only=True) == {1: 1, 4: 1}
    update_word(db, first, WordUpdate(is_favorite=False))
    assert first.knowledge_level == 1 and [topic.id for topic in first.topics] == before_topics
    assert db.scalar(select(func.count()).select_from(WordProgressEvent)) == 0
    configure_language(db, "es")
    assert [word.id for word in get_all_words(db, favorites_only=True)] == [foreign.id]


def test_multi_language_export_keeps_importable_language_metadata_and_does_not_change_scope(db):
    for language, term, favorite in [
        ("en", "hello", True),
        ("es", "hola", True),
        ("es", "adiós", False),
    ]:
        db.add(
            Word(
                term=term,
                language=language,
                is_favorite=favorite,
                topics=[Topic(name=term, slug=term, language=language)],
                translation_items=[WordTranslation(position=0, language="ru", value="привет")],
            )
        )
    db.commit()
    configure_language(db, "en")
    app = FastAPI()
    app.include_router(words_router)
    app.dependency_overrides[get_db] = lambda: db
    with TestClient(app) as client:
        response = client.get(
            "/api/words/export/xlsx?languages=en&languages=es&favorites_only=true"
        )
    assert response.status_code == 200 and response.headers["content-type"] == "application/zip"
    with ZipFile(BytesIO(response.content)) as archive:
        assert set(archive.namelist()) == {"lexora-vocabulary-en.xlsx", "lexora-vocabulary-es.xlsx"}
        for language in ("en", "es"):
            book = load_workbook(BytesIO(archive.read(f"lexora-vocabulary-{language}.xlsx")))
            metadata = dict(book["_Lexora_Language"].values)
            assert metadata["language"] == language
            assert book["Words"].max_row == 2
    assert learning_language(db) == "en"
