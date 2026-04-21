from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.features.words.workbook import cells as workbook_cells
from app.features.words.workbook import format as workbook_format
from app.features.words.workbook import import_rules as workbook_import_rules
from app.features.words.workbook import service as workbook_service


def test_normalize_countability_canonicalizes_known_values() -> None:
    assert workbook_cells._normalize_countability("countable") == "Countable"
    assert workbook_cells._normalize_countability(" UnCountable ") == "Uncountable"
    assert workbook_cells._normalize_countability("both") == "Both"
    assert workbook_cells._normalize_countability("plural") == "Plural"
    assert workbook_cells._normalize_countability("Collective") == "Collective"
    assert workbook_cells._normalize_countability("") is None


def test_countability_validation_range_covers_all_export_values() -> None:
    workbook = Workbook()
    ws = workbook.active
    workbook_format._create_lists_sheet(workbook)
    workbook_format._add_validations(workbook, ws, 5)

    countability_values = [
        workbook[workbook_format.LISTS_SHEET_NAME].cell(row=row, column=2).value
        for row in range(1, len(workbook_format.COUNTABILITY_VALUES) + 1)
    ]
    assert countability_values == workbook_format.COUNTABILITY_VALUES

    countability_validation = [
        dv for dv in ws.data_validations.dataValidation
        if str(dv.sqref) == "F2:F5"
    ][0]
    assert countability_validation.formula1 == (
        f"={workbook_format.LISTS_SHEET_NAME}!$B$1:$B${len(workbook_format.COUNTABILITY_VALUES)}"
    )


def test_countability_for_export_uses_canonical_value() -> None:
    word = SimpleNamespace(countability="uncountable")

    assert workbook_service._countability_for_export(word) == "Uncountable"


def test_meta_sheet_preserves_full_topic_name_when_sheet_title_is_truncated() -> None:
    full_topic_name = "Public Transport City Navigation Long Original Name"
    sheet_title = workbook_format._safe_sheet_title(full_topic_name, set())
    workbook = Workbook()

    workbook_format._create_meta_sheet(workbook, [(sheet_title, 28, full_topic_name)])

    meta = workbook[workbook_format.META_SHEET_NAME]
    assert len(sheet_title) == 31
    assert sheet_title != full_topic_name
    assert meta.cell(row=2, column=1).value == sheet_title
    assert meta.cell(row=2, column=3).value == full_topic_name


def test_existing_word_duplicate_check_skips_unchanged_word_id_row() -> None:
    existing = SimpleNamespace(term="Dog")

    assert workbook_import_rules._should_validate_existing_word_duplicate(
        existing,
        " dog ",
        topic_was_missing=False,
    ) is False


def test_existing_word_duplicate_check_runs_for_changed_term_or_new_topic() -> None:
    existing = SimpleNamespace(term="Dog")

    assert workbook_import_rules._should_validate_existing_word_duplicate(
        existing,
        "hound",
        topic_was_missing=False,
    ) is True
    assert workbook_import_rules._should_validate_existing_word_duplicate(
        existing,
        "dog",
        topic_was_missing=True,
    ) is True


def test_find_existing_word_rejects_word_id_on_wrong_topic(monkeypatch) -> None:
    word = SimpleNamespace(id=42, deleted_at=None, topics=[SimpleNamespace(id=1)])
    monkeypatch.setattr(
        workbook_import_rules,
        "get_word_by_id_including_deleted",
        lambda _db, _word_id: word,
    )

    with pytest.raises(workbook_format.InvalidWorkbookError) as exc_info:
        workbook_import_rules._find_existing_word(
            SimpleNamespace(), topic_id=2, word_id=42, term="dog"
        )

    assert "wrong topic sheet" in str(exc_info.value)


def test_find_existing_word_allows_word_id_on_existing_topic(monkeypatch) -> None:
    word = SimpleNamespace(id=42, deleted_at=None, topics=[SimpleNamespace(id=2)])
    monkeypatch.setattr(
        workbook_import_rules,
        "get_word_by_id_including_deleted",
        lambda _db, _word_id: word,
    )

    assert workbook_import_rules._find_existing_word(
        SimpleNamespace(),
        topic_id=2,
        word_id=42,
        term="dog",
    ) is word


def test_validate_countability_rejects_unknown_value() -> None:
    with pytest.raises(workbook_format.InvalidWorkbookError) as exc_info:
        workbook_cells._validate_countability("sometimes", "Sheet", 7)

    assert "countability must be blank or one of" in str(exc_info.value)


def test_validate_part_of_speech_rejects_unknown_value() -> None:
    with pytest.raises(workbook_format.InvalidWorkbookError) as exc_info:
        workbook_import_rules._validate_part_of_speech("article", "Sheet", 7)

    assert "part of speech must be blank or one of" in str(exc_info.value)


def test_validate_max_length_rejects_too_long_value() -> None:
    with pytest.raises(workbook_format.InvalidWorkbookError) as exc_info:
        workbook_cells._validate_max_length("abcd", 3, "word", "Sheet", 4)

    assert "word must be at most 3 characters" in str(exc_info.value)


def test_import_rejects_oversized_workbook() -> None:
    content = b"0" * (workbook_format.MAX_WORKBOOK_BYTES + 1)

    with pytest.raises(workbook_format.InvalidWorkbookError) as exc_info:
        workbook_service.import_words_workbook(SimpleNamespace(), content)

    assert "Workbook is too large" in str(exc_info.value)


def test_get_topic_rejects_unknown_topic() -> None:
    db = SimpleNamespace(
        scalar=lambda _stmt: None,
    )

    with pytest.raises(workbook_format.InvalidWorkbookError) as exc_info:
        workbook_import_rules._get_topic(db, "Renamed Topic")

    assert "unknown topic" in str(exc_info.value)
