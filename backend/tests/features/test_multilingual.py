"""Exercise separate learning libraries through real API transactions."""

from io import BytesIO
from uuid import uuid4

import pytest
from app.features.languages.catalog import configure_language
from app.features.offline.router import router as offline_router
from app.features.smart_review.router import router as review_router
from app.features.words.agent_router import router as bulk_router
from app.shared.config import settings
from app.shared.deps import get_db, verify_csrf, verify_session
from fastapi import Depends, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from tests.features.test_crud_workbook_flow import build_test_app, topic, word


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(settings, "app_password", "multilingual-test")
    monkeypatch.setattr(settings, "secret_key", "multilingual-test-secret")
    monkeypatch.setattr(settings, "cookie_secure", False)
    monkeypatch.setattr(settings, "api_key", "multilingual-test-api-key")
    app, engine = build_test_app()

    def scoped_db(request: Request):
        with Session(engine, expire_on_commit=False) as db:
            configure_language(db, request.headers.get("X-Lexora-Language", "en"), request.headers.get("X-Lexora-Translation-Language", "ru"))
            yield db

    app.dependency_overrides[get_db] = scoped_db
    for router in (offline_router, review_router):
        app.include_router(router, dependencies=[Depends(verify_session), Depends(verify_csrf)])
    app.include_router(bulk_router)
    with TestClient(app) as result:
        login = result.post("/auth/login", json={"password": "multilingual-test"})
        result.headers["X-CSRF-Token"] = login.json()["csrf_token"]
        yield result
    engine.dispose()


@pytest.mark.parametrize("language,term,name", [
    ("zh", "你好", "问候"), ("es", "hola", "Saludos"), ("de", "hallo", "Grüße"),
    ("tr", "merhaba", "Selamlaşma"), ("tt", "сәлам", "Сәламләшү"),
    ("ar", "مرحبا", "التحيات"), ("fr", "bonjour", "Salutations"), ("pt", "olá", "Saudações"),
])
def test_libraries_keep_words_topics_progress_and_ids_separate(client, language, term, name):
    english = topic(client, "English")
    original = word(client, [english["id"]], knowledge_level=4)
    client.headers["X-Lexora-Language"] = language
    group = topic(client, name)
    added = word(client, [group["id"]], term=term, knowledge_level=1, transliteration="test")
    assert added["language"] == language and group["language"] == language
    assert client.get("/api/words").json()["total"] == 1
    assert client.get("/api/topics").json()[0]["id"] == group["id"]
    assert client.get("/api/words/summary").json() == {"total": 1, "level_counts": {"1": 1}}
    assert client.get(f'/api/words/{original["id"]}').status_code == 404
    assert client.get("/api/words", params={"topic_id": english["id"]}).json()["total"] == 0
    assert client.post("/api/words", json={"term": "foreign", "topic_ids": [english["id"]], "translation_entries": ["x"]}).status_code == 400
    assert client.post("/api/topics", json={"name": "bad parent", "parent_topic_id": english["id"]}).status_code == 400
    client.headers["X-Lexora-Language"] = "en"
    assert client.get("/api/words").json()["words"][0]["id"] == original["id"]
    assert client.get("/api/words/summary").json() == {"total": 1, "level_counts": {"4": 1}}


def test_same_topic_and_spelling_allowed_in_other_language(client):
    en = topic(client, "General")
    word(client, [en["id"]], term="gift")
    client.headers["X-Lexora-Language"] = "de"
    de = topic(client, "General")
    added = word(client, [de["id"]], term="gift", part_of_speech=None)
    assert added["language"] == "de"
    assert client.post("/api/words", json={"term": "gift", "topic_ids": [de["id"]], "translation_entries": ["яд"]}).status_code == 409
    assert client.post("/api/words", json={"language": "en", "term": "other", "topic_ids": [de["id"]], "translation_entries": ["x"]}).status_code == 400


def test_translation_and_explanation_languages_do_not_overwrite_each_other(client):
    client.headers["X-Lexora-Language"] = "es"
    group = topic(client, "Food")
    added = word(client, [group["id"]], term="pan", definition="Продукт из муки", notes="Русская заметка")
    response = client.put(f'/api/words/{added["id"]}', json={"translation_language": "en", "translation_entries": ["bread"], "definition": "A baked food", "notes": "English note"})
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["translation_groups"] == {"ru": ["чек", "квитанция"], "en": ["bread"]}
    assert data["definition"] == "Продукт из муки"
    assert data["explanation_groups"]["en"] == {"definition": "A baked food", "notes": "English note"}
    client.put(f'/api/words/{added["id"]}', json={"translation_entries": ["хлеб"], "knowledge_level": 3})
    data = client.get(f'/api/words/{added["id"]}').json()
    assert data["translation_groups"] == {"ru": ["хлеб"], "en": ["bread"]}
    assert data["knowledge_level"] == 3


def test_daily_mix_refresh_leaves_other_language_queue_active(client):
    en = topic(client, "English")
    first = word(client, [en["id"]], knowledge_level=1)
    mix = client.get("/api/smart-review")
    assert mix.status_code == 200, mix.text
    original = mix.json()
    client.headers["X-Lexora-Language"] = "fr"
    fr = topic(client, "French")
    second = word(client, [fr["id"]], term="bonjour", knowledge_level=1)
    other = client.post("/api/smart-review/refresh").json()
    assert [item["word_id"] for item in other["items"]] == [second["id"]]
    client.headers["X-Lexora-Language"] = "en"
    current = client.get("/api/smart-review").json()
    assert current["id"] == original["id"]
    assert current["items"][0]["word_id"] == first["id"]


def test_workbook_scope_and_saved_preview_cannot_cross_languages(client):
    en = topic(client, "General")
    word(client, [en["id"]])
    client.headers["X-Lexora-Language"] = "de"
    de = topic(client, "General")
    target = word(client, [de["id"]], term="Brot")
    exported = client.get("/api/words/export/xlsx", params={"layout": "flat"})
    workbook = load_workbook(BytesIO(exported.content))
    assert workbook["Words"].max_row == 2
    assert workbook["Words"].cell(2, 12).value == target["id"]
    workbook["Words"].cell(2, 3, "хлеб")
    stream = BytesIO()
    workbook.save(stream)
    response = client.post("/api/words/import/xlsx/preview", files={"file": ("words.xlsx", stream.getvalue())}, data={"mode": "content"})
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["errors"] == []
    payload = {"batch_id": preview["batch_id"], "plan_hash": preview["plan_hash"]}
    client.headers["X-Lexora-Language"] = "en"
    assert client.post("/api/words/import/xlsx/apply", json=payload).status_code == 409
    client.headers["X-Lexora-Language"] = "de"
    applied = client.post("/api/words/import/xlsx/apply", json=payload)
    assert applied.status_code == 200, applied.text
    assert client.get(f'/api/words/{target["id"]}').json()["translation_entries"] == ["хлеб"]


def test_offline_updates_use_own_library_and_preserve_other_translations(client):
    client.headers["X-Lexora-Language"] = "tr"
    group = topic(client, "General")
    saved = word(client, [group["id"]], term="ekmek")
    change = {"id": str(uuid4()), "entity": "word", "action": "update", "entity_id": saved["id"], "before": saved,
        "data": {"translation_entries": ["bread"], "translation_language": "en", "definition": "A food"}}
    response = client.post("/api/offline/changes", json=change)
    assert response.status_code == 200, response.text
    assert response.json()["item"]["translation_groups"]["ru"] == saved["translation_entries"]
    client.headers["X-Lexora-Language"] = "en"
    change["id"] = str(uuid4())
    assert client.post("/api/offline/changes", json=change).status_code == 409
    assert client.get("/api/offline/snapshot").json()["words"] == []
    assert client.get("/api/offline/snapshot", headers={"X-Lexora-All-Languages": "true"}).json()["words"][0]["language"] == "tr"


def test_translation_search_and_workbook_metadata_use_the_selected_language(client):
    group = topic(client, "General")
    saved = word(client, [group["id"]], term="bread", definition="Продукт из муки")
    client.put(f'/api/words/{saved["id"]}', json={"translation_entries": ["bread synonym"], "translation_language": "en", "definition": "Baked food"})
    workbook = client.get("/api/words/export/xlsx", params={"layout": "flat"})
    client.headers["X-Lexora-Translation-Language"] = "en"
    assert client.get("/api/words", params={"search": "квитанция"}).json()["total"] == 0
    assert client.get("/api/words", params={"search": "Baked"}).json()["total"] == 1
    preview = client.post("/api/words/import/xlsx/preview", files={"file": ("russian.xlsx", workbook.content)})
    assert any("another learning or translation language" in issue["message"] for issue in preview.json()["errors"])
    assert preview.json()["batch_id"] is None
    english_export = client.get("/api/words/export/xlsx", params={"layout": "flat"})
    book = load_workbook(BytesIO(english_export.content))
    assert book["Words"].cell(2, 3).value == "bread synonym"
    assert book["Words"].cell(2, 8).value == "Baked food"
    book["Words"].cell(2, 3, "A loaf")
    stream = BytesIO()
    book.save(stream)
    preview = client.post("/api/words/import/xlsx/preview", files={"file": ("english.xlsx", stream.getvalue())}, data={"mode": "content"}).json()
    assert preview["errors"] == []
    result = client.post("/api/words/import/xlsx/apply", json={key: preview[key] for key in ("batch_id", "plan_hash")})
    assert result.status_code == 200, result.text
    data = client.get(f'/api/words/{saved["id"]}').json()
    assert data["translation_groups"] == {"ru": ["чек", "квитанция"], "en": ["A loaf"]}
    assert data["definition"] == "Продукт из муки" and data["knowledge_level"] == 4


def test_statistics_usage_and_trash_are_scoped(client):
    en = topic(client, "General")
    first = word(client, [en["id"]], term="one", knowledge_level=1)
    client.put(f'/api/words/{first["id"]}', json={"knowledge_level": 4})
    usage = {"event_key": str(uuid4()), "session_key": "multilingual-session", "route": "/", "active_seconds": 30}
    assert client.post("/api/stats/usage", json=usage).status_code == 204
    client.headers["X-Lexora-Language"] = "fr"
    fr = topic(client, "General")
    second = word(client, [fr["id"]], term="un", knowledge_level=1)
    stats = client.get("/api/stats").json()
    assert stats["usage_summary"]["total_active_seconds"] == 0
    assert stats["efficiency_summary"]["total_review_events"] == 0
    assert [item["id"] for item in stats["topics"]] == [fr["id"]]
    client.delete(f'/api/words/{second["id"]}')
    assert len(client.get("/api/trash/words").json()) == 1
    client.headers["X-Lexora-Language"] = "en"
    assert client.get("/api/trash/words").json() == []
    old = word(client, [en["id"]], term="discard")
    client.delete(f'/api/words/{old["id"]}')
    assert len(client.get("/api/trash/words").json()) == 1
    assert client.delete("/api/trash/purge", params={"force": True}).status_code == 204
    assert client.get("/api/trash/words").json() == []
    client.headers["X-Lexora-Language"] = "fr"
    assert len(client.get("/api/trash/words").json()) == 1


def test_json_bulk_import_defaults_to_selected_library_and_translation_target(client):
    client.headers["X-Api-Key"] = settings.api_key
    client.headers["X-Lexora-Language"] = "zh"
    client.headers["X-Lexora-Translation-Language"] = "en"
    response = client.post("/api/words/bulk", json={"topic_name": "问候", "words": [{"term": "你好", "translation_entries": ["hello"], "definition": "A greeting", "transliteration": "nǐ hǎo"}]})
    assert response.status_code == 201, response.text
    saved = client.get("/api/words").json()["words"][0]
    assert saved["language"] == "zh" and saved["transliteration"] == "nǐ hǎo"
    assert saved["translation_groups"] == {"en": ["hello"]}
    assert saved["explanation_groups"]["en"]["definition"] == "A greeting"
    response = client.post("/api/words/bulk", json={"topic_name": "wrong", "words": [{"term": "hello", "language": "en"}]})
    assert response.status_code == 400


def test_offline_legacy_pinyin_edit_and_multilingual_delete_conflict(client):
    client.headers["X-Lexora-Language"] = "zh"
    group = topic(client, "General")
    saved = word(client, [group["id"]], term="你好")
    legacy = {key: value for key, value in saved.items() if key != "transliteration"}
    change = {"id": str(uuid4()), "entity": "word", "action": "update", "entity_id": saved["id"], "before": legacy, "data": {"transliteration": "nǐ hǎo"}}
    response = client.post("/api/offline/changes", json=change)
    assert response.status_code == 200, response.text
    snapshot = response.json()["item"]
    assert client.put(f'/api/words/{saved["id"]}', json={"translation_language": "en", "translation_entries": ["hello"], "definition": "A greeting"}).status_code == 200
    stale_delete = {"id": str(uuid4()), "entity": "word", "action": "delete", "entity_id": saved["id"], "before": snapshot}
    response = client.post("/api/offline/changes", json=stale_delete)
    assert response.status_code == 409, response.text
    assert client.get(f'/api/words/{saved["id"]}').json()["translation_groups"]["en"] == ["hello"]
