"""Exercise real routers, transactions and XLSX files against a disposable database."""

from io import BytesIO

import pytest
from fastapi import Depends, FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.features.auth.router import router as auth_router
from app.features.health.router import router as health_router
from app.features.smart_review import model as _review_model  # noqa: F401
from app.features.stats.router import router as stats_router
from app.features.topics.router import router as topics_router
from app.features.trash.router import router as trash_router
from app.features.words.constants import PART_OF_SPEECH_VALUES
from app.features.words.model import PartsOfSpeech
from app.features.words.router import router as words_router
from app.shared.config import settings
from app.shared.db import Base
from app.shared.deps import get_db, verify_csrf, verify_session


def build_test_app():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        poolclass=StaticPool,
        connect_args={"check_same_thread": False},
    )
    event.listen(
        engine, "connect", lambda connection, _: connection.execute("PRAGMA foreign_keys=ON")
    )
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        db.add_all(PartsOfSpeech(name=name) for name in PART_OF_SPEECH_VALUES)
        db.commit()

    def test_db():
        with Session(engine, expire_on_commit=False) as db:
            yield db

    app = FastAPI()
    app.dependency_overrides[get_db] = test_db
    app.include_router(auth_router)
    app.include_router(health_router)
    for router in (topics_router, words_router, trash_router, stats_router):
        app.include_router(router, dependencies=[Depends(verify_session), Depends(verify_csrf)])
    return app, engine


@pytest.fixture
def api_client(monkeypatch):
    monkeypatch.setattr(settings, "app_password", "crud-test-password")
    monkeypatch.setattr(settings, "secret_key", "crud-test-secret")
    monkeypatch.setattr(settings, "cookie_secure", False)
    app, engine = build_test_app()
    with TestClient(app, raise_server_exceptions=False) as client:
        login = client.post("/auth/login", json={"password": "crud-test-password"})
        assert login.status_code == 200
        client.headers["X-CSRF-Token"] = login.json()["csrf_token"]
        yield client
    engine.dispose()


def topic(client, name, parent=None):
    response = client.post("/api/topics", json={"name": name, "parent_topic_id": parent})
    assert response.status_code == 201, response.text
    return response.json()


def word(client, topic_ids, term="receipt", **fields):
    response = client.post(
        "/api/words",
        json={
            "term": term,
            "topic_ids": topic_ids,
            "translation_entries": ["чек", "квитанция"],
            "part_of_speech": "noun",
            "knowledge_level": 4,
            **fields,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_batch_ratings_record_history_once_per_actual_change(api_client):
    client = api_client
    group = topic(client, "History")
    first = word(client, [group["id"]], term="first", knowledge_level=1)
    second = word(client, [group["id"]], term="second", knowledge_level=3)
    payload = {"word_ids": [first["id"], second["id"]], "knowledge_level": 4}
    for _ in range(2):
        assert client.patch("/api/words/batch", json=payload).json() == {"updated": 2}

    stats = client.get("/api/stats").json()
    assert stats["efficiency_summary"]["total_review_events"] == 2
    assert stats["daily_activity"][0]["reviewed"] == 2
    assert stats["daily_activity"][0]["improved"] == 2
    assert stats["retention_summary"]["strong_words"] == 2


@pytest.mark.parametrize("deleted", [False, True])
def test_batch_rejects_unavailable_topics_without_saving_ratings(api_client, deleted):
    client = api_client
    group = topic(client, "Keep")
    entry = word(client, [group["id"]], knowledge_level=1)
    unavailable = 999999
    if deleted:
        unavailable = topic(client, "Deleted")["id"]
        assert client.delete(f"/api/topics/{unavailable}").status_code == 204
    response = client.patch("/api/words/batch", json={
        "word_ids": [entry["id"]], "knowledge_level": 4, "add_topic_ids": [unavailable],
    })
    assert response.status_code == 400
    assert "Topics not found" in response.json()["detail"]
    saved = client.get(f"/api/words/{entry['id']}").json()
    assert saved["knowledge_level"] == 1
    assert saved["topic_ids"] == [group["id"]]


def test_batch_rejects_last_topic_removal_atomically_and_allows_replacement(api_client):
    client = api_client
    shared = topic(client, "Shared")
    other = topic(client, "Other")
    first = word(client, [shared["id"], other["id"]], term="first", knowledge_level=1)
    second = word(client, [shared["id"]], term="second", knowledge_level=1)
    payload = {"word_ids": [first["id"], second["id"]], "knowledge_level": 4,
        "remove_topic_ids": [shared["id"]]}
    response = client.patch("/api/words/batch", json=payload)
    assert response.status_code == 400
    assert "at least one active topic" in response.json()["detail"]
    for entry in (first, second):
        saved = client.get(f"/api/words/{entry['id']}").json()
        assert saved["knowledge_level"] == 1
        assert set(saved["topic_ids"]) == set(entry["topic_ids"])
    assert client.get("/api/stats").json()["efficiency_summary"]["total_review_events"] == 0

    payload["add_topic_ids"] = [other["id"]]
    assert client.patch("/api/words/batch", json=payload).json() == {"updated": 2}
    for entry in (first, second):
        assert client.get(f"/api/words/{entry['id']}").json()["topic_ids"] == [other["id"]]


def preview(client, book, mode="content"):
    stream = BytesIO()
    book.save(stream)
    return client.post(
        "/api/words/import/xlsx/preview",
        data={"mode": mode},
        files={
            "file": (
                "words.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def apply_preview(client, result):
    return client.post(
        "/api/words/import/xlsx/apply",
        json={"batch_id": result["batch_id"], "plan_hash": result["plan_hash"]},
    )


def upload(client, book):
    response = preview(client, book)
    if response.status_code != 200 or response.json()["errors"]:
        return response
    return apply_preview(client, response.json())


def exported(client):
    response = client.get("/api/words/export/xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(BytesIO(response.content))


def test_topic_and_word_lifecycle_with_shared_membership(api_client):
    client = api_client
    parent = topic(client, "Shopping")
    child = topic(client, "Receipts", parent["id"])
    other = topic(client, "Money")
    receipt = word(client, [child["id"], other["id"]])
    exclusive = word(client, [child["id"]], "invoice")
    assert client.post("/api/topics", json={"name": "Shopping"}).status_code == 409
    assert client.delete(f"/api/topics/{parent['id']}").status_code == 409
    assert (
        client.put(f"/api/topics/{parent['id']}", json={"parent_topic_id": child["id"]}).status_code
        == 400
    )
    renamed = client.put(
        f"/api/topics/{other['id']}", json={"name": "Payments", "description": "Updated topic"}
    )
    assert renamed.status_code == 200 and renamed.json()["slug"] == "payments"
    updated = client.put(
        f"/api/words/{receipt['id']}",
        json={
            "term": "payment receipt",
            "translation_entries": ["чек об оплате", "квитанция об оплате"],
            "example_entries": ["EN: Keep the receipt.\nRU: Сохраните чек."],
        },
    )
    assert updated.status_code == 200 and updated.json()["knowledge_level"] == 4
    assert updated.json()["translation_entries"] == ["чек об оплате", "квитанция об оплате"]
    assert client.delete(f"/api/topics/{child['id']}").status_code == 204
    assert client.get(f"/api/words/{receipt['id']}").status_code == 200
    assert client.get(f"/api/words/{exclusive['id']}").status_code == 404
    restored = client.post(f"/api/trash/topics/{child['id']}/restore?restore_words=true")
    assert restored.status_code == 200
    assert client.get(f"/api/words/{exclusive['id']}").status_code == 200
    assert client.delete(f"/api/words/{receipt['id']}").status_code == 204
    assert client.get(f"/api/words/{receipt['id']}").status_code == 404
    assert client.post(f"/api/trash/words/{receipt['id']}/restore").status_code == 200
    assert client.get(f"/api/words/{receipt['id']}").json()["knowledge_level"] == 4


def test_export_import_preserves_mixed_examples_and_shared_word_ids(api_client):
    client = api_client
    first, second = topic(client, "Shopping"), topic(client, "Money")
    examples = ["Keep the receipt.", "EN: The receipt is here.\nRU: Чек здесь."]
    saved = word(client, [first["id"], second["id"]], example_entries=examples)
    book = exported(client)
    response = upload(client, book)
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 0
    current = client.get(f"/api/words/{saved['id']}").json()
    assert current["example_entries"] == examples
    assert current["topic_ids"] == saved["topic_ids"]
    assert current["knowledge_level"] == saved["knowledge_level"]
    assert upload(client, exported(client)).status_code == 200


@pytest.mark.parametrize(("column", "value"), [(1, 4.9), (1, True), (12, 1.5), (12, True)])
def test_import_rejects_non_integer_ids_and_levels_without_writes(api_client, column, value):
    client = api_client
    group = topic(client, "Shopping")
    saved = word(client, [group["id"]])
    book = exported(client)
    book["Shopping"].cell(2, column, value)
    book["Shopping"].cell(2, 3, "must not be saved")
    response = upload(client, book)
    assert response.status_code == 200 and response.json()["errors"], response.text
    assert client.get(f"/api/words/{saved['id']}").json() == saved


def test_duplicate_import_is_a_validation_error_and_rolls_back_earlier_rows(api_client):
    client = api_client
    first = topic(client, "Shopping")
    topic(client, "Money")
    saved = word(client, [first["id"]])
    book = exported(client)
    book["Shopping"].cell(2, 3, "must not be saved")
    book["Money"].cell(2, 2, "receipt")
    book["Money"].cell(2, 3, "duplicate from a different sheet")
    response = upload(client, book)
    assert response.status_code == 200 and response.json()["errors"], response.text
    assert "different content" in response.json()["errors"][0]["message"]
    assert client.get(f"/api/words/{saved['id']}").json() == saved
    assert client.get("/api/words").json()["total"] == 1


def test_workbook_can_add_new_words_and_edit_existing_words(api_client):
    client = api_client
    group = topic(client, "Shopping")
    saved = word(client, [group["id"]])
    book = exported(client)
    sheet = book["Shopping"]
    sheet.cell(2, 3, "новый перевод\nдругая версия")
    sheet.cell(3, 1, 2)
    sheet.cell(3, 2, "invoice")
    sheet.cell(3, 3, "счёт")
    sheet.cell(3, 7, "noun")
    response = upload(client, book)
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 1
    assert response.json()["updated"] == 1
    assert client.get(f"/api/words/{saved['id']}").json()["translation_entries"] == [
        "новый перевод",
        "другая версия",
    ]
    assert client.get("/api/words").json()["total"] == 2


@pytest.mark.parametrize("examples", ["RU: Без английского.", "EN:\nRU: Перевод.", "EN: One.\nRU:"])
def test_import_still_rejects_malformed_bilingual_examples(api_client, examples):
    client = api_client
    group = topic(client, "Shopping")
    saved = word(client, [group["id"]])
    book = exported(client)
    book["Shopping"].cell(2, 5, examples)
    response = preview(client, book)
    assert response.status_code == 200 and response.json()["errors"]
    assert client.get(f"/api/words/{saved['id']}").json() == saved


def flat_book(rows, headers=("Word", "Translations", "Part of speech", "Topic IDs", "Word ID")):
    book = Workbook()
    book.active.title = "Words"
    book.active.append(list(headers))
    for row in rows:
        book.active.append(row)
    return book


def test_preview_is_non_mutating_apply_is_durable_and_retryable(api_client):
    client = api_client
    group = topic(client, "Shopping")
    book = flat_book([["receipt", "чек", "noun", str(group["id"]), None]])
    plan = preview(client, book, "add").json()
    assert not plan["errors"] and plan["summary"]["created"] == 1
    assert client.get("/api/words").json()["total"] == 0
    applied = apply_preview(client, plan)
    assert applied.status_code == 200, applied.text
    assert applied.json()["created"] == 1
    repeated = apply_preview(client, plan)
    assert repeated.json()["replayed"] is True
    assert repeated.json()["created_word_ids"] == applied.json()["created_word_ids"]
    assert client.get("/api/words").json()["total"] == 1
    recovery = client.get(f"/api/words/import/xlsx/{plan['batch_id']}/recovery").json()
    from hashlib import sha256

    from app.features.words.workbook.sync import canonical_json

    assert sha256(canonical_json(recovery["payload"])).hexdigest() == recovery["sha256"]
    assert recovery["payload"]["created_word_ids"] == applied.json()["created_word_ids"]
    again = preview(client, book, "add").json()
    assert again["summary"]["created"] == again["summary"]["updated"] == 0
    assert again["summary"]["unchanged"] == 1


def test_blank_id_keeps_different_parts_of_speech_separate(api_client):
    client = api_client
    group = topic(client, "Everyday")
    noun = word(client, [group["id"]], "watch", translation_entries=["часы"])
    plan = preview(
        client, flat_book([["watch", "смотреть", "verb", group["id"], None]]), "add"
    ).json()
    assert not plan["errors"]
    applied = apply_preview(client, plan)
    assert applied.status_code == 200, applied.text
    assert client.get(f"/api/words/{noun['id']}").json() == noun
    assert client.get("/api/words").json()["total"] == 2
    assert client.delete(f'/api/words/{noun["id"]}').status_code == 204
    assert client.post(f'/api/trash/words/{noun["id"]}/restore').status_code == 200
    assert client.get(f'/api/words/{noun["id"]}').json()["part_of_speech"] == "noun"


def test_blank_id_cannot_guess_a_different_sense(api_client):
    group = topic(api_client, "Money")
    saved = word(api_client, [group["id"]], "bank", translation_entries=["банк"])
    plan = preview(api_client, flat_book([["bank", "берег", "noun", group["id"], None]])).json()
    assert plan["errors"] and plan["batch_id"] is None
    assert api_client.get(f"/api/words/{saved['id']}").json() == saved


def test_import_preserves_blank_fields_and_newer_progress(api_client):
    client = api_client
    group = topic(client, "Shopping")
    saved = word(client, [group["id"]], notes="My note", definition="Proof of payment.")
    book = exported(client)
    sheet = book["Shopping"]
    for column in (1, 8, 11):
        sheet.cell(2, column).value = None
    sheet.cell(2, 3, "чек об оплате")
    plan = preview(client, book).json()
    assert client.put(f"/api/words/{saved['id']}", json={"knowledge_level": 3}).status_code == 200
    applied = apply_preview(client, plan)
    assert applied.status_code == 200, applied.text
    current = client.get(f"/api/words/{saved['id']}").json()
    assert current["knowledge_level"] == 3
    assert current["notes"] == saved["notes"] and current["definition"] == saved["definition"]
    assert current["translation_entries"] == ["чек об оплате"]
    recovery = client.get(f"/api/words/import/xlsx/{plan['batch_id']}/recovery").json()
    assert (
        recovery["payload"]["changes"][0]["before"]["translations"] == saved["translation_entries"]
    )


def test_content_change_after_preview_rejects_entire_apply(api_client):
    client = api_client
    group = topic(client, "Shopping")
    saved = word(client, [group["id"]])
    book = exported(client)
    book["Shopping"].cell(2, 3, "запланированный перевод")
    plan = preview(client, book).json()
    client.put(f"/api/words/{saved['id']}", json={"translation_entries": ["новый ручной перевод"]})
    response = apply_preview(client, plan)
    assert response.status_code == 409
    assert api_client.get(f"/api/words/{saved['id']}").json()["translation_entries"] == [
        "новый ручной перевод"
    ]


def test_new_shared_word_and_new_topic_are_created_once(api_client):
    client = api_client
    group = topic(client, "Shopping")
    book = flat_book(
        [["receipt", "чек", "noun", group["id"], None, "Payments"]],
        headers=("Word", "Translations", "Part of speech", "Topic IDs", "Word ID", "Topics"),
    )
    sheet = book.create_sheet("Topics")
    sheet.append(["Topic", "Parent ID", "Description"])
    sheet.append(["Payments", group["id"], "Payments topic"])
    plan = preview(client, book, "add").json()
    assert not plan["errors"] and plan["summary"]["topics_created"] == 1
    result = apply_preview(client, plan).json()
    saved = client.get(f"/api/words/{result['created_word_ids'][0]}").json()
    assert set(saved["topic_ids"]) == {group["id"], result["created_topic_ids"][0]}
    again = preview(client, book, "add").json()
    assert not again["errors"] and again["summary"]["unchanged"] == 1
    assert again["summary"]["topics_created"] == 0


def test_collects_all_errors_and_does_not_apply_valid_rows(api_client):
    group = topic(api_client, "Shopping")
    book = flat_book(
        [
            ["receipt", "чек", "noun", group["id"], None],
            [None, "осиротевший перевод", "noun", group["id"], None],
            ["invoice", "счёт", "not-pos", group["id"], None],
        ]
    )
    plan = preview(api_client, book).json()
    assert len(plan["errors"]) == 2 and plan["batch_id"] is None
    assert {item["row"] for item in plan["errors"]} == {3, 4}
    assert api_client.get("/api/words").json()["total"] == 0


def test_flat_export_has_one_row_per_word_and_reimports_unchanged(api_client):
    client = api_client
    first, second = topic(client, "Shopping"), topic(client, "Money")
    word(client, [first["id"], second["id"]])
    response = client.get("/api/words/export/xlsx?layout=flat")
    book = load_workbook(BytesIO(response.content))
    assert book["Words"].max_row == 2
    assert len(str(book["Words"].cell(2, 13).value).splitlines()) == 2
    plan = preview(client, book).json()
    assert not plan["errors"] and plan["summary"]["unchanged"] == 1
    assert plan["summary"]["updated"] == 0
    template = client.get("/api/words/import/xlsx/template")
    assert template.status_code == 200
    assert load_workbook(BytesIO(template.content))["Words"].max_row == 1


def test_large_import_uses_bounded_selects(api_client):
    from app.features.words.workbook.import_models import WorkbookImportBatch
    from app.shared.deps import get_db

    client = api_client
    group = topic(client, "Shopping")
    dependency = client.app.dependency_overrides[get_db]()
    db = next(dependency)
    selects = []
    engine = db.get_bind()

    def track(_connection, _cursor, statement, *_args):
        if statement.lstrip().upper().startswith("SELECT"):
            selects.append(statement)

    event.listen(engine, "before_cursor_execute", track)
    try:
        book = flat_book(
            [[f"term {index}", "перевод", "noun", group["id"], None] for index in range(7000)]
        )
        plan = preview(client, book, "add").json()
        assert not plan["errors"] and plan["summary"]["created"] == 7000
        result = apply_preview(client, plan)
        assert result.status_code == 200, result.text
        assert result.json()["created"] == 7000
        assert len(selects) < 100, len(selects)
        assert db.get(WorkbookImportBatch, plan["batch_id"]).result["created"] == 7000
    finally:
        event.remove(engine, "before_cursor_execute", track)
        dependency.close()


def test_failed_apply_rolls_back_content_and_can_retry_same_receipt(api_client, monkeypatch):
    from app.features.words.workbook import import_batches

    client = api_client
    group = topic(client, "Shopping")
    saved = word(client, [group["id"]])
    book = exported(client)
    book["Shopping"].cell(2, 3, "новое значение")
    plan = preview(client, book).json()

    def fail_after_old_entries_were_removed(**_):
        raise RuntimeError("Simulated failure during insert")

    with monkeypatch.context() as patch:
        patch.setattr(import_batches, "WordTranslation", fail_after_old_entries_were_removed)
        assert apply_preview(client, plan).status_code == 500
    assert client.get(f"/api/words/{saved['id']}").json() == saved
    retried = apply_preview(client, plan)
    assert retried.status_code == 200 and retried.json()["updated"] == 1
    assert not retried.json()["replayed"]


def test_word_added_after_preview_prevents_duplicate_create(api_client):
    group = topic(api_client, "Shopping")
    plan = preview(
        api_client, flat_book([["receipt", "чек", "noun", group["id"], None]]), "add"
    ).json()
    word(api_client, [group["id"]], translation_entries=["чек"])
    assert apply_preview(api_client, plan).status_code == 409
    assert api_client.get("/api/words").json()["total"] == 1


@pytest.mark.parametrize(
    "headers,rows",
    [
        (["Word", "Term", "Translations", "Topic IDs"], [["one", "two", "перевод", "1"]]),
        (["Word", "Translations", "Topic IDs"], [['=UPPER("one")', "перевод", "1"]]),
    ],
)
def test_duplicate_headers_and_formulas_are_reported(api_client, headers, rows):
    topic(api_client, "Shopping")
    plan = preview(api_client, flat_book(rows, headers)).json()
    assert plan["errors"] and plan["batch_id"] is None
    assert api_client.get("/api/words").json()["total"] == 0


def test_receipt_migration_upgrades_and_downgrades_without_touching_vocabulary():
    import importlib.util
    from pathlib import Path

    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from sqlalchemy import inspect, text

    path = (
        Path(__file__).resolve().parents[2]
        / "alembic/versions/20261014_0010_workbook_import_batches.py"
    )
    spec = importlib.util.spec_from_file_location("receipt_migration", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE words (id INTEGER PRIMARY KEY, term TEXT)"))
        connection.execute(text("INSERT INTO words VALUES (1, 'preserved')"))
        migration.op = Operations(MigrationContext.configure(connection))
        migration.upgrade()
        assert "workbook_import_batches" in inspect(connection).get_table_names()
        migration.downgrade()
        assert "workbook_import_batches" not in inspect(connection).get_table_names()
        assert connection.execute(text("SELECT term FROM words")).scalar() == "preserved"
    engine.dispose()
